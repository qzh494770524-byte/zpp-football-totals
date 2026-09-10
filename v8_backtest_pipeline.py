# -*- coding: utf-8 -*-
"""
V8.0 大小球回测流水线(可复用版)

用法:
    python v8_backtest_pipeline.py "<Nowgoal多庄水位比分.xlsx路径>" [--no-fetch] [--out <输出目录>]

输入文件要求: 与 Nowgoal_2026-09-09_多庄水位比分.xlsx 相同的sheet结构
    - 全部赛事与比分
    - 多庄亚洲盘
    - 各庄大小球与欧赔

这个脚本做四件事:
  1. 读取三个sheet,对每场比赛算出"盘口漂移模型"方向(线体漂移+水位漂移+欧赔平局代理+让球深浅,
     权重沿用第一轮回测验证过的基线: line20/water15/euro10/handicap10)。
  2. 对同一批比赛,批量抓取 live11.nowgoal26.com/match/h2h-{比赛ID} 页面(自动限速,已抓过的会缓存跳过),
     解析"Standings"板块拿到两队 总/主场/客场/近6场 的真实进球/失球数据,算出"进球模型"期望总进球。
  3. 合并两个模型:
       - 方向一致 -> 历史回测里这类场次命中率没有优势(约39%),标记为"无明显edge"
       - 方向冲突 -> 历史回测里以进球模型为准命中率明显更高(66.7% vs 33.3%,但样本仅15场,需持续验证),
         标记为"优先关注",并按进球模型数据样本量给出信心高低
       - 进球数据不足(球队近期比赛太少) -> 退回只看盘口漂移模型,标记"仅盘口信号"
  4. 如果比赛已经完场(有终场比分),自动核对预测是否命中,输出整体/分组命中率统计;
     如果比赛还没开始,只输出预测方向,供你之后回来核对。

重要声明: "冲突时信进球模型"这条规则,目前只在一批89场里的15场冲突样本上验证过,
按V8.0原文档"十四、赛后复盘"的要求,至少要积累20-30场独立样本才能真正确认,
在此之前请把这个脚本的输出当作"待验证假设"而不是可下注的结论。
"""
import argparse
import json
import os
import re
import statistics as st
import time
import urllib.error
import urllib.request
from collections import Counter, defaultdict

import openpyxl

HEADERS = {'User-Agent': 'Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36'}
H2H_URL = 'https://live11.nowgoal26.com/match/h2h-{}'
REQUEST_DELAY_SEC = 1.5


# ---------------------------------------------------------------------------
# 第一步: 读取Excel, 计算盘口漂移模型
# ---------------------------------------------------------------------------

def sign(x, eps):
    if x > eps:
        return 1
    if x < -eps:
        return -1
    return 0


def implied_prob_pair(a, b):
    ia, ib = 1.0 / (1.0 + a), 1.0 / (1.0 + b)
    s = ia + ib
    return ia / s, ib / s


def load_workbook_data(xlsx_path, min_companies=2):
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)

    matches = {}
    for r in wb['全部赛事与比分'].iter_rows(min_row=2, values_only=True):
        (mid, league, kickoff, home, away, status, final_score, cur_score,
         n_ah, n_verified, list_complete, score_time) = r
        if not n_ah or n_ah < min_companies:
            continue  # 跳过覆盖公司数不足 min_companies 家的赛事
        total_goals, hs, as_ = None, None, None
        if final_score:
            try:
                hs, as_ = map(int, final_score.split('-'))
                total_goals = hs + as_
            except Exception:
                pass
        matches[mid] = dict(league=league, home=home, away=away, status=status,
                             final_score=final_score, total_goals=total_goals)

    ah_by_match = defaultdict(list)
    for r in wb['多庄亚洲盘'].iter_rows(min_row=2, values_only=True):
        (mid, league, kickoff, home, away, status, cid, cname,
         open_hw, open_line, open_aw, last_hw, last_line, last_aw, last_time,
         live_hw, live_line, live_aw, verify_status, fh, fa, fscore,
         fill_status, snap_time, src, note) = r
        if mid not in matches or open_line is None or last_line is None:
            continue
        ah_by_match[mid].append(dict(open_hw=open_hw, open_line=open_line,
                                      last_hw=last_hw, last_line=last_line))

    ou_by_match = defaultdict(list)
    eu_by_match = defaultdict(list)
    for r in wb['各庄大小球与欧赔'].iter_rows(min_row=2, values_only=True):
        (mid, league, home, away, cid, cname, ptype,
         open_a, open_mid, open_b, live_a, live_mid, live_b, fscore) = r
        if mid not in matches:
            continue
        if ptype == '大小球(香港水位)':
            if None in (open_a, open_mid, open_b, live_a, live_mid, live_b):
                continue
            ou_by_match[mid].append(dict(open_over=open_a, open_line=open_mid, open_under=open_b,
                                          live_over=live_a, live_line=live_mid, live_under=live_b))
        elif ptype == '胜平负(欧洲赔率)':
            if None in (open_a, open_mid, open_b, live_a, live_mid, live_b):
                continue
            eu_by_match[mid].append(dict(open_draw=open_mid, live_draw=live_mid))

    return matches, ah_by_match, ou_by_match, eu_by_match


def compute_odds_drift_signal(mid, ou_rows, ah_rows, eu_rows):
    """返回 None(数据不足) 或 dict(direction, confidence_prob, consensus_line, support_count, trap_flag)"""
    if len(ou_rows) < 2:
        return None

    line_signals, water_signals, consensus_live_lines = [], [], []
    for row in ou_rows:
        line_signals.append(sign(row['live_line'] - row['open_line'], 0.001))
        p_over_open, _ = implied_prob_pair(row['open_over'], row['open_under'])
        p_over_live, _ = implied_prob_pair(row['live_over'], row['live_under'])
        water_signals.append(sign(p_over_live - p_over_open, 0.01))
        consensus_live_lines.append(row['live_line'])

    line_sig = st.mean(line_signals)
    water_sig = st.mean(water_signals)
    consensus_line = st.median(consensus_live_lines)

    dirs = [d for d in (1 if x > 0 else (-1 if x < 0 else 0) for x in line_signals) if d != 0]
    consistency_dir = Counter(dirs).most_common(1)[0][0] if dirs else 0

    euro_sig = st.mean([sign(r['live_draw'] - r['open_draw'], 0.01) for r in eu_rows]) if eu_rows else 0.0

    handicap_sig, trap_flag = 0.0, False
    if ah_rows:
        deltas, trap_votes = [], 0
        for row in ah_rows:
            d = row['last_line'] - row['open_line']
            deltas.append(sign(d, 0.01))
            hw_delta = ((row['last_hw'] - row['open_hw'])
                        if (row['last_hw'] is not None and row['open_hw'] is not None) else 0)
            if d > 0 and hw_delta > 0.01:
                trap_votes += 1
            if d < 0 and hw_delta < -0.01:
                trap_votes += 1
        handicap_sig = st.mean(deltas) if deltas else 0.0
        trap_flag = trap_votes >= max(1, len(ah_rows) // 2)

    weights = {'line': 20, 'water': 15, 'euro': 10, 'handicap': 10}
    total_w = sum(weights.values())
    sig = {'line': line_sig, 'water': water_sig, 'euro': euro_sig, 'handicap': handicap_sig}
    composite = sum(sig[k] * w for k, w in weights.items()) / total_w
    composite += (0.05 if consistency_dir > 0 else (-0.05 if consistency_dir < 0 else 0)) * (5 / total_w) * 20
    composite = max(-1, min(1, composite))

    prob_over = max(20, min(80, 50 + composite * 20))
    direction = 'over' if prob_over >= 50 else 'under'
    confidence_prob = prob_over if direction == 'over' else 100 - prob_over

    dir_target = 1 if direction == 'over' else -1
    support_count = sum(1 for s in (sign(line_sig, .001), sign(water_sig, .001),
                                     sign(euro_sig, .001), sign(handicap_sig, .001)) if s == dir_target)
    if support_count < 3:
        confidence_prob = min(confidence_prob, 57.0)

    return dict(direction=direction, confidence_prob=round(confidence_prob, 1),
                consensus_line=consensus_line, support_count=support_count, trap_flag=trap_flag)


# ---------------------------------------------------------------------------
# 第二步: 抓取 + 解析 h2h 页面 -> 进球模型
# ---------------------------------------------------------------------------

def fetch_h2h_html(mid, cache_dir):
    os.makedirs(cache_dir, exist_ok=True)
    cache_path = os.path.join(cache_dir, f"{mid}.html")
    if os.path.exists(cache_path):
        with open(cache_path, encoding='utf-8') as f:
            return f.read()
    req = urllib.request.Request(H2H_URL.format(mid), headers=HEADERS)
    try:
        with urllib.request.urlopen(req, timeout=20) as resp:
            html = resp.read().decode('utf-8', errors='replace')
    except (urllib.error.URLError, TimeoutError) as e:
        print(f"  [警告] 抓取 {mid} 失败: {e}")
        return None
    with open(cache_path, 'w', encoding='utf-8') as f:
        f.write(html)
    time.sleep(REQUEST_DELAY_SEC)
    return html


def _strip_tags(s):
    return re.sub(r'<[^>]+>', '', s).strip()


def _parse_standings_table(table_html):
    m = re.search(r"openFbTeam\((\d+)\)'>\s*(?:\[([^\]]+)\])?\s*&nbsp;&nbsp;([^<]+?)\s*</a>", table_html)
    if not m:
        return None
    team_id, league_rank, team_name = m.group(1), m.group(2), m.group(3).strip()
    ft_match = re.search(r"<th >FT</th>.*?(?=<th class='ht-desc'>HT</th>|$)", table_html, re.S)
    ft_html = ft_match.group(0) if ft_match else ''
    stats = {}
    for row in re.findall(r"<tr align='center'>(.*?)</tr>", ft_html, re.S):
        tds = re.findall(r"<td[^>]*>(.*?)</td>", row, re.S)
        if len(tds) < 9:
            continue
        label = _strip_tags(tds[0])
        if label not in ('Total', 'Home', 'Away', 'Last 6'):
            continue
        try:
            matches_, win, draw, lose, scored, conceded, pts, rank, rate = [_strip_tags(x) for x in tds[1:10]]
            stats[label] = dict(matches=int(matches_), scored=int(scored), conceded=int(conceded),
                                 rank=(int(rank) if rank.strip() else None))
        except ValueError:
            continue
    return dict(team_id=team_id, league_rank_label=league_rank, team_name=team_name, stats=stats)


def parse_h2h_standings(html):
    idx = html.find('>Standings<')
    if idx == -1:
        return None, None
    segment = html[idx:idx + 12000]
    tables = re.findall(r'<table.*?</table>', segment, re.S)
    if len(tables) < 2:
        return None, None
    return _parse_standings_table(tables[0]), _parse_standings_table(tables[1])


def compute_expected_goals(home, away):
    hs, as_ = home['stats'], away['stats']

    def pick_rate(stats, venue, key, min_n=2):
        v = stats.get(venue)
        if v and v.get('matches', 0) >= min_n:
            return v[key] / v['matches'], v['matches']
        l6 = stats.get('Last 6')
        if l6 and l6.get('matches', 0) >= 2:
            return l6[key] / l6['matches'], l6['matches']
        t = stats.get('Total')
        if t and t.get('matches', 0) >= 2:
            return t[key] / t['matches'], t['matches']
        return None, 0

    home_scored_home, n1 = pick_rate(hs, 'Home', 'scored')
    away_conceded_away, n2 = pick_rate(as_, 'Away', 'conceded')
    away_scored_away, n3 = pick_rate(as_, 'Away', 'scored')
    home_conceded_home, n4 = pick_rate(hs, 'Home', 'conceded')

    if None in (home_scored_home, away_conceded_away, away_scored_away, home_conceded_home):
        return None

    lam_home = (home_scored_home + away_conceded_away) / 2
    lam_away = (away_scored_away + home_conceded_home) / 2
    return lam_home + lam_away, min(n1, n2, n3, n4)


# ---------------------------------------------------------------------------
# 主流程
# ---------------------------------------------------------------------------

def main():
    ap = argparse.ArgumentParser(description='V8.0 大小球回测流水线(可复用)')
    ap.add_argument('xlsx', help='Nowgoal多庄水位比分.xlsx 文件路径')
    ap.add_argument('--no-fetch', action='store_true', help='跳过抓取h2h页面,只跑盘口漂移模型')
    ap.add_argument('--out', default=None, help='输出目录(默认: xlsx同目录下的 v8_pipeline_out)')
    ap.add_argument('--min-companies', type=int, default=2, help='只保留覆盖公司数>=此值的赛事(默认2)')
    args = ap.parse_args()

    xlsx_path = args.xlsx
    out_dir = args.out or os.path.join(os.path.dirname(os.path.abspath(xlsx_path)), 'v8_pipeline_out')
    cache_dir = os.path.join(out_dir, 'h2h_cache')
    os.makedirs(out_dir, exist_ok=True)

    print(f"读取: {xlsx_path}")
    matches, ah_by_match, ou_by_match, eu_by_match = load_workbook_data(xlsx_path, min_companies=args.min_companies)
    print(f"总赛事数: {len(matches)}")

    per_match = {}
    for mid, m in matches.items():
        odds_sig = compute_odds_drift_signal(mid, ou_by_match.get(mid, []),
                                              ah_by_match.get(mid, []), eu_by_match.get(mid, []))
        per_match[mid] = dict(match=m, odds_signal=odds_sig, goal_signal=None)

    n_odds_ok = sum(1 for v in per_match.values() if v['odds_signal'])
    print(f"有盘口漂移信号的赛事数(>=2家公司大小球数据): {n_odds_ok}")

    if not args.no_fetch:
        candidate_ids = [mid for mid, v in per_match.items() if v['odds_signal']]
        print(f"\n开始抓取/读取缓存 {len(candidate_ids)} 场比赛的h2h页面 (间隔{REQUEST_DELAY_SEC}s)...")
        for i, mid in enumerate(candidate_ids, 1):
            html = fetch_h2h_html(mid, cache_dir)
            if html is None:
                continue
            home, away = parse_h2h_standings(html)
            if home is None or away is None:
                continue
            eg = compute_expected_goals(home, away)
            if eg is None:
                continue
            expected_total, min_sample = eg
            line = per_match[mid]['odds_signal']['consensus_line']
            diff = expected_total - line
            direction = 'over' if diff > 0 else 'under'
            prob = max(20, min(80, 50 + diff * 20))
            confidence_prob = prob if direction == 'over' else 100 - prob
            per_match[mid]['goal_signal'] = dict(
                expected_total=round(expected_total, 2), diff=round(diff, 2),
                direction=direction, confidence_prob=round(confidence_prob, 1), min_sample=min_sample,
                home_team=home['team_name'], away_team=away['team_name'])
            if i % 10 == 0 or i == len(candidate_ids):
                print(f"  进度 {i}/{len(candidate_ids)}")
        n_goal_ok = sum(1 for v in per_match.values() if v['goal_signal'])
        print(f"成功算出进球模型的赛事数: {n_goal_ok}")

    # ---------------- 生成最终建议 ----------------
    final_rows = []
    for mid, v in per_match.items():
        m, odds_sig, goal_sig = v['match'], v['odds_signal'], v['goal_signal']
        if odds_sig is None:
            continue
        if goal_sig is None:
            recommend_dir = odds_sig['direction']
            confidence_prob = odds_sig['confidence_prob']
            tag = '仅盘口信号(进球数据不足)'
        elif goal_sig['direction'] == odds_sig['direction']:
            recommend_dir = odds_sig['direction']
            confidence_prob = (odds_sig['confidence_prob'] + goal_sig['confidence_prob']) / 2
            tag = '两模型一致(历史回测无明显edge,约39%命中率,谨慎参考)'
        else:
            recommend_dir = goal_sig['direction']
            confidence_prob = goal_sig['confidence_prob']
            conf_note = '高' if goal_sig['min_sample'] >= 3 else '低(球队数据样本<3场)'
            tag = f'两模型冲突->优先关注,以进球模型为准(历史66.7% vs 33.3%,数据信心{conf_note})'

        row = dict(mid=mid, league=m['league'], home=m['home'], away=m['away'],
                   status=m['status'], final_score=m['final_score'],
                   consensus_line=odds_sig['consensus_line'],
                   odds_direction=odds_sig['direction'], odds_confidence=odds_sig['confidence_prob'],
                   goal_direction=(goal_sig['direction'] if goal_sig else None),
                   goal_expected_total=(goal_sig['expected_total'] if goal_sig else None),
                   recommend_direction=recommend_dir, recommend_confidence=round(confidence_prob, 1),
                   tag=tag)

        if m['total_goals'] is not None:
            actual = ('over' if m['total_goals'] > odds_sig['consensus_line']
                      else 'under' if m['total_goals'] < odds_sig['consensus_line'] else 'push')
            row['actual'] = actual
            row['hit'] = (recommend_dir == actual) if actual != 'push' else None
        else:
            row['actual'] = None
            row['hit'] = None

        final_rows.append(row)

    out_json = os.path.join(out_dir, 'predictions.json')
    with open(out_json, 'w', encoding='utf-8') as f:
        json.dump(final_rows, f, ensure_ascii=False, indent=2)
    print(f"\n逐场结果已保存: {out_json}")

    # ---------------- 汇总报告(只对已完场且非push的场次统计) ----------------
    scored = [r for r in final_rows if r['hit'] is not None]
    pending = [r for r in final_rows if r['status'] != '完场']
    print(f"\n未开赛/进行中(仅预测,无法核对): {len(pending)} 场")
    print(f"已完场可核对: {len(scored)} 场")
    if scored:
        hits = sum(1 for r in scored if r['hit'])
        print(f"整体命中率: {hits}/{len(scored)} = {hits/len(scored)*100:.1f}%")
        for group_name, cond in [
            ('两模型一致', lambda r: '一致' in r['tag']),
            ('两模型冲突(优先关注)', lambda r: '冲突' in r['tag']),
            ('仅盘口信号', lambda r: '仅盘口' in r['tag']),
        ]:
            grp = [r for r in scored if cond(r)]
            if not grp:
                continue
            h = sum(1 for r in grp if r['hit'])
            print(f"  {group_name}: {h}/{len(grp)} = {h/len(grp)*100:.1f}%")

    print("\n=== 冲突场次(优先关注)明细 ===")
    for r in final_rows:
        if '冲突' not in r['tag']:
            continue
        status_mark = ''
        if r['hit'] is True:
            status_mark = '[OK]'
        elif r['hit'] is False:
            status_mark = '[X]'
        elif r['status'] != '完场':
            status_mark = '[待开赛/进行中]'
        print(f"{status_mark} {r['league'][:20]:20s} {r['home'][:14]:14s} vs {r['away'][:14]:14s} "
              f"盘口{r['consensus_line']} 建议{r['recommend_direction']}({r['recommend_confidence']}%) "
              f"实际={r['actual']}")

    print(f"\n完整结果与建议见: {out_json}")


if __name__ == '__main__':
    main()
