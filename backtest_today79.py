"""Fixed-rule retrospective O/U test for the 79 newly settled fixtures on Sep 10."""
from collections import Counter, defaultdict
from datetime import datetime
import hashlib
import json
import os
from pathlib import Path
import statistics
import sys
from urllib.parse import parse_qs, urlparse

# The backtest does not need BLAS workers; keep optional NumPy imports lightweight.
os.environ['OPENBLAS_NUM_THREADS'] = '1'
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

import analyze_totals as totals
import nowgoal_collect as core
import nowgoal_multi as multi

OUT = core.HERE / 'analysis' / 'backtest_79_20260910'
BOOKS = [8, 3, 31, 50]
DAY = '2026-09-10'
END = '2026-09-10T08:33:02+08:00'


def save(name, value):
    (OUT / name).write_text(json.dumps(value, ensure_ascii=False, indent=2), encoding='utf-8')


def prepare():
    OUT.mkdir(parents=True, exist_ok=True)
    if (OUT/'cohort.json').exists():
        return
    log = (core.HERE/'history_watcher.log').read_bytes()
    current = {m['id']: m for m in totals.read_json(core.HERE/'multi_collection_2026-09-09.json')}
    baseline, latest, stamp = None, None, None
    groups = defaultdict(list)
    summaries = []
    for lineno, line in enumerate(log.decode('utf-8').splitlines(), 1):
        try:
            item = json.loads(line)
        except ValueError:
            continue
        if 'historical_start' in item and 'updated_at' in item and 'finished_matches' in item:
            stamp = item['updated_at']
            if stamp < DAY:
                baseline = item
            elif stamp <= END:
                summaries.append(item)
                latest = item
        if stamp and stamp.startswith(DAY) and stamp <= END and item.get('url'):
            q = parse_qs(urlparse(item['url']).query)
            if q.get('type') == ['14'] and q.get('t') == ['1'] and q.get('s') == ['-1'] and int(q['id'][0]) in current:
                groups[stamp].append({'id': int(q['id'][0]), 'log_line': lineno, 'cycle': stamp})
    assert baseline and latest and baseline['comparison_pending_matches'] == 0
    previous = baseline['finished_matches']
    for item in summaries:
        added = item['finished_matches'] - previous
        unique = {r['id'] for r in groups.get(item['updated_at'], [])}
        assert len(unique) == added, f"Cohort mismatch at {item['updated_at']}: {len(unique)} vs {added}"
        previous = item['finished_matches']
    calls = {r['id']: r for group in groups.values() for r in group}
    assert len(calls) == latest['finished_matches']-baseline['finished_matches'] == 79
    cohort = []
    truths = []
    for mid, call in calls.items():
        m = current[mid]
        cohort.append({**call, **{k:m[k] for k in ('league','home','away','kickoff')}})
        truths.append({k:m.get(k) for k in ('id','home_score','away_score','state','explain','league')})
    save('cohort.json', cohort)
    save('outcomes.json', truths)
    save('cohort_audit.json', {'baseline_time':baseline['updated_at'], 'baseline_finished':baseline['finished_matches'],
                              'end_time':latest['updated_at'], 'end_finished':latest['finished_matches'],
                              'log_sha256':hashlib.sha256(log).hexdigest(), 'per_cycle':{k:len(v) for k,v in groups.items()},
                              'reconstruction':'Each cycle final-state refresh IDs reconcile exactly with newly finished counts; 79 unique IDs.'})
    save('fixed_rule.json', {
        'specified_at':core.stamp(), 'kind':'retrospective fixed-rule test; not a real-time prediction record',
        'company_priority':BOOKS, 'minimum_same_line_books':3,
        'quote_selection':'Latest type 1/2 non-closed full-match OU history record strictly before scheduled kickoff; history downloaded after first finished refresh.',
        'line_selection':'Line quoted by most eligible books; ties resolved by fixed company priority.',
        'direction':'For each book compute (1/(1+overHK))/(1/(1+overHK)+1/(1+underHK)); median >0.5 selects over, <0.5 selects under, equal skips.',
        'settlement_quote':'First available bookmaker in fixed priority at the chosen line; do not choose best odds after outcomes.',
        'missing_policy':'No forecast if fewer than 3 same-line bookmakers; no outcome grading if regulation-time score is unclear.',
        'metrics':'Full win, half win, push, half loss, full loss; positive-result share excluding pushes; net units and ROI per 1 unit wager.',
        'baselines':'Always-over and always-under using precisely the same settled fixtures, lines and bookmaker quotes.',
        'tuning':'No fitting or parameter search against these 79 results.'})


def collect():
    prepare()
    core.OFFLINE = False; core.SOURCE_LIMITED = False
    data = []
    cohort = totals.read_json(OUT/'cohort.json')
    for index, m in enumerate(cohort, 1):
        mid = m['id']; entry = {**m, 'quotes':[], 'errors':[]}
        try:
            comp = core.get_json(core.odds_url(mid,-1),f'comp_{mid}.json',refresh=False)
            offered = {c['cid']:c for c in comp['Data']['mixodds']}
            for cid in BOOKS:
                company = offered.get(cid, {})
                if not (totals.quote(company.get('ou',{}).get('f')) or totals.quote(company.get('ou',{}).get('l'))):
                    continue
                path = core.ROOT/f'history_{mid}_{cid}.json'
                refresh = not path.exists() or path.stat().st_mtime < datetime.fromisoformat(m['cycle']).timestamp()
                try:
                    history = core.get_json(core.history_url(mid,cid), path.name, refresh=refresh)
                    cutoff = datetime.fromisoformat(m['kickoff']).timestamp()
                    records = [r for r in history.get('Data',{}).get('ou',[]) if r.get('type') in (1,2)
                               and not r.get('close') and totals.quote(r.get('odds')) and r.get('mt',float('inf')) < cutoff]
                    if records:
                        record=max(records,key=lambda r:r['mt'])
                        entry['quotes'].append({'cid':cid,'company':multi.name(cid),**totals.quote(record['odds']),
                                                 'record_time':datetime.fromtimestamp(record['mt'],core.TZ).isoformat(),
                                                 'history_file':str(path),'history_sha256':hashlib.sha256(path.read_bytes()).hexdigest()})
                except Exception as error:
                    entry['errors'].append(f'{multi.name(cid)}: {error}')
                    if core.SOURCE_LIMITED:
                        break
        except Exception as error:
            entry['errors'].append(str(error))
        data.append(entry)
        save('pregame_inputs.json',data)
        if index%10==0:
            print(json.dumps({'collected_fixtures':index,'total':len(cohort),'source_limited':core.SOURCE_LIMITED}),flush=True)
        if core.SOURCE_LIMITED:
            # Preserve all 79 cohort rows, explicitly recording those still uncollected.
            data.extend({**remaining,'quotes':[],'errors':['Source restriction; pregame histories pending']} for remaining in cohort[index:])
            save('pregame_inputs.json',data)
            break


def forecast(quotes):
    grouped=defaultdict(list)
    for q in quotes:
        if q['cid'] in BOOKS:
            grouped[q['line']].append(q)
    if not grouped:
        return {'side':None,'reason':'无可核验赛前大小球记录'}
    line, group=min(grouped.items(),key=lambda kv:(-len(kv[1]),min(BOOKS.index(q['cid']) for q in kv[1])))
    if len({q['cid'] for q in group})<3:
        return {'side':None,'reason':'同一盘口不足3家有效记录','line':line,'books':len(group)}
    indicators=[(1/(1+q['over']))/(1/(1+q['over'])+1/(1+q['under'])) for q in group]
    indicator=statistics.median(indicators)
    if abs(indicator-.5)<1e-12:
        return {'side':None,'reason':'综合水位指标恰好相等','line':line,'books':len(group)}
    anchor=min(group,key=lambda q:BOOKS.index(q['cid']))
    return {'side':'over' if indicator>.5 else 'under','line':line,'books':len(group),'indicator':indicator,
            'anchor_company':anchor['company'],'over_water':anchor['over'],'under_water':anchor['under'],
            'record_time':anchor['record_time'],'reason':'按固定多庄赛前规则'}


def evaluate():
    inputs=totals.read_json(OUT/'pregame_inputs.json')
    # Forecasts are materialized before the outcome file is read.
    predictions=[{**{k:m[k] for k in ('id','league','home','away','kickoff','cycle')},**forecast(m['quotes'])} for m in inputs]
    save('locked_predictions.json',predictions)
    truths={r['id']:r for r in totals.read_json(OUT/'outcomes.json')}
    results=[]
    for p in predictions:
        truth=truths[p['id']]
        r=dict(p,home_score=truth['home_score'],away_score=truth['away_score'])
        reason=totals.exclusion(truth)
        r['settled']=False
        if not p['side']:
            r['status']=p['reason']
        elif reason:
            r['status']='不结算：'+reason
        elif truth['state']!=-1 or truth['home_score'] is None or truth['away_score'] is None:
            r['status']='不结算：终场比分缺失'
        else:
            total=truth['home_score']+truth['away_score']
            result=totals.settle(total,p['line'],p[p['side']+'_water'],p['side'])
            r.update(settled=True,total=total,result=result['result'],profit=result['profit'],status='已结算')
            r['always_over']=totals.settle(total,p['line'],p['over_water'],'over')['profit']
            r['always_under']=totals.settle(total,p['line'],p['under_water'],'under')['profit']
        results.append(r)
    save('results.json',results)
    settled=[r for r in results if r['settled']]
    count=Counter(r['result'] for r in settled)
    positive=count['全赢']+count['半赢']; negative=count['全输']+count['半输']
    summary={'cohort':len(results),'predictions':sum(r['side'] is not None for r in results),'settled':len(settled),
             'skipped_no_forecast':sum(r['side'] is None for r in results),'skipped_score_scope':sum(r['side'] is not None and not r['settled'] for r in results),
             'outcome_counts':dict(count),'positive_matches':positive,'negative_matches':negative,
             'positive_rate_ex_push':positive/(positive+negative) if positive+negative else None,
             'positive_rate_ci':totals.wilson(positive,positive+negative),
             'positive_rate_all_settled':positive/len(settled) if settled else None,
             'net_units':sum(r['profit'] for r in settled),'roi':statistics.mean(r['profit'] for r in settled) if settled else None,
             'always_over_roi':statistics.mean(r['always_over'] for r in settled) if settled else None,
             'always_under_roi':statistics.mean(r['always_under'] for r in settled) if settled else None,
             'skip_reasons':dict(Counter(r['status'] for r in results if not r['settled'])),
             'completed_at':core.stamp(), 'claims':'Retrospective fixed-rule evaluation, not an advance prediction or validated profitable strategy.'}
    save('summary.json',summary)
    wb=Workbook();ws=wb.active;ws.title='79场逐场回测'
    ws.append(['比赛ID','回填时间','联赛','主队','客队','预测','盘口','同盘口公司数','多庄水位指标','结算公司','大球水位','小球水位','赛前记录时间','终场比分','结算结果','净单位','状态'])
    for r in results:
        ws.append([r['id'],r['cycle'],r['league'],r['home'],r['away'],{'over':'大球','under':'小球'}.get(r['side'],'跳过'),
                   r.get('line'),r.get('books'),r.get('indicator'),r.get('anchor_company'),r.get('over_water'),r.get('under_water'),
                   r.get('record_time'),f"{r['home_score']}-{r['away_score']}",r.get('result'),r.get('profit'),r['status']])
    hs=wb.create_sheet('赛前输入凭据');hs.append(['比赛ID','公司','盘口','大球水位','小球水位','赛前记录时间','来源文件'])
    for m in inputs:
        for q in m['quotes']:
            hs.append([m['id'],q['company'],q['line'],q['over'],q['under'],q['record_time'],q['history_file']])
    notes=wb.create_sheet('统计与规则');notes.append(['项目','值'])
    for k,v in summary.items():notes.append([k,json.dumps(v,ensure_ascii=False) if isinstance(v,(dict,list,tuple)) else v])
    for k,v in totals.read_json(OUT/'fixed_rule.json').items():notes.append([k,json.dumps(v,ensure_ascii=False) if isinstance(v,(dict,list)) else v])
    notes.append(['胜率口径','(全赢+半赢)/(全赢+半赢+全输+半输)；另给含走盘的盈利场次比例。半赢不等于全赢收益。'])
    notes.append(['限制','这79场已有赛果。本次没有训练或挑选最佳参数，但事后回测仍不能替代真正赛前锁定预测的验证。'])
    notes.append(['水位指标','归一化报价仅用于选方向，不是已经校准的真实进球概率。'])
    notes.append(['队列核验','每轮s=-1的刷新ID数与日志完场增量逐轮一致，合计79个唯一ID。详见cohort_audit.json。'])
    for sheet in wb:
        sheet.freeze_panes='A2';sheet.auto_filter.ref=sheet.dimensions
        for c in sheet[1]:c.font=Font(bold=True,color='FFFFFF');c.fill=PatternFill('solid',fgColor='17365D')
        for col in sheet.columns:
            sheet.column_dimensions[get_column_letter(col[0].column)].width=min(44,max(15,max(len(str(c.value or '')) for c in list(col)[:80])+2))
        for row in sheet:
            for cell in row:
                if cell.data_type=='f':cell.data_type='s'
    wb.save(OUT/'新增79场_大小球固定规则回测.xlsx')
    pct=totals.pct
    report=['# 今日新增79场：大小球固定规则回测','',
            f"79场中，{summary['predictions']}场形成预测，{summary['settled']}场可核对结算，{summary['skipped_no_forecast']}场资料不足或指标相等而跳过，{summary['skipped_score_scope']}场因赛果口径跳过。",'',
            f"全赢{count['全赢']}、半赢{count['半赢']}、走盘{count['走盘']}、半输{count['半输']}、全输{count['全输']}。",
            f"剔除走盘后的盈利场次率：{pct(summary['positive_rate_ex_push'])}；95%区间：{pct(summary['positive_rate_ci'][0])}～{pct(summary['positive_rate_ci'][1])}。",
            f"每场1单位，净{summary['net_units']:+.2f}单位，模拟ROI {pct(summary['roi'])}。",'',
            f"相同样本、相同盘口与结算公司：始终选大球的ROI {pct(summary['always_over_roi'])}，始终选小球的ROI {pct(summary['always_under_roi'])}。基线仅用于比较，不按较高者重新选择规则。",'',
            '规则：固定Bet365、Crown、Sbobet、1xBet；仅采用有时间核验的赛前大小球末条，至少3家同盘口。对双方水位换算的归一化报价指标取中位数，大于0.5选大、小于0.5选小、相等跳过。结算公司按事先固定顺序选择。',
            '该指标没有经过胜率校准，不能把数值直接称为真实获胜概率。整数和季度盘分别计算走盘、半赢、半输。',
            '这些比赛已有赛果；这是固定规则的历史回测，不是事先发布的预测，也未证明可持续盈利。缺数据的比赛不计入命中率分母。',
            'cohort_audit.json记录79场ID的日志核对依据；fixed_rule.json记录先行确定的规则；locked_predictions.json在评估函数读取赛果文件前写出；Excel包含全部79场及赛前输入。']
    (OUT/'回测结论.md').write_text('\n'.join(report),encoding='utf-8')
    print(json.dumps(summary,ensure_ascii=True),flush=True)


if __name__=='__main__':
    if '--evaluate-only' not in sys.argv:
        collect()
    evaluate()
