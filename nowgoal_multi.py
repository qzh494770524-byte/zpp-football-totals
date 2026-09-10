"""Multi-bookmaker collection, historical pre-match verification and score backfill."""
import argparse
from collections import Counter
from datetime import datetime, timedelta
import json
from pathlib import Path
import time

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

import nowgoal_collect as core

NAMES = {8: 'Bet365', 3: 'Crown', 31: 'Sbobet', 50: '1xBet', 17: 'M88',
         24: '12Bet', 42: '18Bet', 12: 'Easybet', 1: 'Macauslot',
         4: 'Ladbrokes', 14: 'Vcbet', 19: 'Interwetten'}
STARTED = {-1, 1, 2, 3, 4, 5}


def name(cid):
    return NAMES.get(cid, f'公司ID {cid}')


def read_success(filename):
    path = core.ROOT / filename
    if path.exists():
        value = json.loads(path.read_text(encoding='utf-8'))
        if value.get('ErrCode') == 0:
            return value
    return None


def eligible(company):
    return any(core.complete(company.get('ah', {}).get(k)) for k in ('f', 'l'))


def company_history(match, cid):
    history = match.get('multi_history', {}).get(str(cid), {})
    return core.prematch({'kickoff': match['kickoff'], 'history': history})


def import_cached(matches):
    for match in matches:
        mid = match['id']
        comp = read_success(f'comp_{mid}.json')
        if comp:
            match['companies'] = comp['Data']['mixodds']
            match['comparison_complete'] = True
            match['comparison_at'] = datetime.fromtimestamp((core.ROOT / f'comp_{mid}.json').stat().st_mtime, core.TZ).isoformat()
        else:
            match.setdefault('comparison_complete', False)
            match.setdefault('comparison_at', match['collected_at'])
        histories = match.setdefault('multi_history', {})
        errors = match.setdefault('multi_errors', {})
        for company in match.get('companies', []):
            cid = company['cid']
            cached = read_success(f'history_{mid}_{cid}.json')
            if cached:
                histories[str(cid)] = cached['Data']
                errors.pop(str(cid), None)
        # The legacy collector's history field belongs only to the selected company.
        selected_cid = match.get('selected', {}).get('cid')
        if selected_cid and match.get('history') and str(selected_cid) not in histories:
            histories[str(selected_cid)] = match['history']


def verified_close(match, cid):
    if str(cid) in match.get('history_refresh_needed', []):
        return None
    records = company_history(match, cid)
    return records[-1] if match['state'] in STARTED and records else None


def score(match):
    if match['state'] == -1 and match.get('home_score') is not None and match.get('away_score') is not None:
        return f"{match['home_score']}-{match['away_score']}"
    return None


def export(matches, day, source_status, scope_description=None, update_command=None):
    wb = Workbook()
    ws = wb.active
    ws.title = '多庄亚洲盘'
    ws.append(['比赛ID', '联赛', '开赛时间(北京时间)', '主队', '客队', '比赛状态', '公司ID', '赔率公司',
               '初盘主水', '初盘盘口(正=主让)', '初盘客水',
               '赛前最后主水', '赛前最后盘口(正=主让)', '赛前最后客水', '赛前最后记录时间(北京时间)',
               '页面即时主水', '页面即时盘口(正=主让)', '页面即时客水', '历史核验状态',
               '终场主队比分', '终场客队比分', '终场比分', '比分回填状态', '水位快照时间', '来源', '备注'])
    roster = wb.create_sheet('全部赛事与比分')
    roster.append(['比赛ID', '联赛', '开赛时间(北京时间)', '主队', '客队', '比赛状态', '终场比分',
                   '当前比分', '有亚洲盘公司数', '已核验临场公司数', '公司列表是否完整', '比分数据时间'])
    hs = wb.create_sheet('各庄赛前历史')
    hs.append(['比赛ID', '公司ID', '赔率公司', '记录时间(北京时间)', '主水', '盘口(正=主让)', '客水', '源站赛前类型'])
    snapshots = wb.create_sheet('各庄大小球与欧赔')
    snapshots.append(['比赛ID', '联赛', '主队', '客队', '公司ID', '赔率公司', '盘口类型',
                      '初盘大球水位或主胜赔率', '初盘盘口或平局赔率', '初盘小球水位或客胜赔率',
                      '页面即时大球水位或主胜赔率', '页面即时盘口或平局赔率', '页面即时小球水位或客胜赔率', '终场比分'])
    counts = Counter()
    coverage = {}
    for match in sorted(matches, key=lambda m: (m['kickoff'], m['id'])):
        companies = [c for c in match.get('companies', []) if eligible(c)]
        final = score(match)
        verified_count = sum(verified_close(match, c['cid']) is not None for c in companies)
        counts['matches'] += 1
        counts['finished_matches'] += final is not None
        counts['pending_matches'] += match['state'] >= 0
        counts['matches_with_multiple_ah_books'] += len(companies) >= 2
        counts['matches_with_ah'] += bool(companies)
        counts['comparison_pending_matches'] += not match.get('comparison_complete') or bool(match.get('comparison_refresh_needed'))
        state_text = core.STATES.get(match['state'], str(match['state']))
        roster.append([match['id'], match['league'], match['kickoff'], match['home'], match['away'], state_text, final,
                       f"{match['home_score']}-{match['away_score']}" if match['state'] in STARTED else None,
                       len(companies), verified_count, '已获取' if match.get('comparison_complete') else '待补采',
                       match.get('scores_at', match['collected_at'])])
        # Retain every fixture, including those for which no bookmaker quote is available.
        for company in companies or [None]:
            cid = company['cid'] if company else None
            ah = company['ah'] if company else {}
            closing = verified_close(match, cid) if cid is not None else None
            records = company_history(match, cid) if cid is not None else []
            hist_fetched = str(cid) in match.get('multi_history', {}) and str(cid) not in match.get('history_refresh_needed', [])
            note = []
            if company is None:
                status = '源站未提供亚洲盘报价' if match.get('comparison_complete') else '公司与水位待补采'
            elif match['state'] == 0:
                status = '未开赛，临场水位待形成'
            elif closing:
                status = '已按本公司赛前历史核验'
            elif hist_fetched:
                status = '历史中无可核验赛前记录'
            elif match['state'] in STARTED:
                status = '本公司赛前历史待补采'
            else:
                status = '异常赛况，待确认'
            if cid is not None and match.get('multi_errors', {}).get(str(cid)):
                note.append(match['multi_errors'][str(cid)])
            if not match.get('comparison_complete'):
                note.append('目前仅有部分公司快照，公司列表待补齐')
            if closing and core.complete(ah.get('l')) and core.trio(closing['odds']) != core.trio(ah['l']):
                note.append('页面即时盘与历史末条不同，临场列采用本公司历史记录')
            ws.append([match['id'], match['league'], match['kickoff'], match['home'], match['away'], state_text,
                       cid, name(cid) if cid is not None else '待补采/无报价', *core.trio(ah.get('f')),
                       *core.trio(closing['odds'] if closing else {}),
                       datetime.fromtimestamp(closing['mt'], core.TZ).isoformat() if closing else None,
                       *core.trio(ah.get('l')), status, match['home_score'] if final else None,
                       match['away_score'] if final else None, final,
                       '已回填' if final else ('待回填' if match['state'] >= 0 else '异常赛况'),
                       match['comparison_at'], f"{core.BASE}/oddscomp/{match['id']}", '; '.join(note)])
            if company is not None:
                item = coverage.setdefault(cid, Counter())
                item['matches'] += 1
                item['initial'] += core.complete(ah.get('f'))
                item['verified'] += closing is not None
                item['history_pending'] += match['state'] in STARTED and not hist_fetched
                item['finished'] += final is not None
                counts['bookmaker_match_rows'] += 1
                counts['initial_ah_rows'] += core.complete(ah.get('f'))
                counts['verified_closing_rows'] += closing is not None
                counts['history_pending_rows'] += match['state'] in STARTED and not hist_fetched
                for record in records:
                    hs.append([match['id'], cid, name(cid), datetime.fromtimestamp(record['mt'], core.TZ).isoformat(),
                               *core.trio(record['odds']), record['type']])
        for company in match.get('companies', []):
            for market, title in [('ou', '大小球(香港水位)'), ('euro', '胜平负(欧洲赔率)')]:
                odds = company.get(market, {})
                if any(core.complete(odds.get(key)) for key in ('f', 'l')):
                    snapshots.append([match['id'], match['league'], match['home'], match['away'], company['cid'],
                                      name(company['cid']), title, *core.trio(odds.get('f')), *core.trio(odds.get('l')), final])
    counts['ah_bookmakers'] = len(coverage)
    cs = wb.create_sheet('公司覆盖统计')
    cs.append(['公司ID', '赔率公司', '有亚洲盘赛事数', '有初盘赛事数', '已核验临场赛事数', '待补赛前历史赛事数', '已回填比分赛事数'])
    for cid, c in sorted(coverage.items(), key=lambda item: -item[1]['matches']):
        cs.append([cid, name(cid), c['matches'], c['initial'], c['verified'], c['history_pending'], c['finished']])
    notes = wb.create_sheet('说明')
    for row in [
        ['项目', '说明'], ['日期范围', scope_description or f'北京时间{day} 00:00至次日00:00，保留所有赛事。'],
        ['多庄口径', '每场比赛×每家有报价的公司一行；公司ID用于关联，绝不将一家公司的历史用于另一家公司。'],
        ['亚洲盘初盘', '各公司接口f字段；水位为香港格式，盘口正数表示主队让球、负数表示主队受让。'],
        ['临场核验', '已开赛/完场：取本公司标记赛前(type=1或2)、未封盘、早于计划开赛的最后一条公开历史记录。'],
        ['时间边界', '源站未提供实际开球时间。若延迟开球，计划时间前的末条不一定是实际开球前的末条。'],
        ['页面即时盘', '各公司l字段，区别于滚球r字段；不将未经历史核验的即时盘直接称为临场盘。'],
        ['缺失值', '未开盘、未开赛、未抓到及无可核验历史均按状态标注；空白不补零。'],
        ['比分', '仅完场填终场比分；同一比赛的终场比分回填到该比赛所有公司的行。'],
        ['其他盘口', '附表保留大小球和欧赔的初盘、页面即时快照；未将其称为已核验的临场水位。'],
        ['数据完整性', source_status], ['采集统计', json.dumps(dict(counts), ensure_ascii=False)],
        ['更新命令', update_command or f'python nowgoal_multi.py --date {day} --watch --interval 1800'], ['更新时间', core.stamp()]]:
        notes.append(row)
    for sheet in wb:
        sheet.freeze_panes = 'A2'
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.font = Font(bold=True, color='FFFFFF')
            cell.fill = PatternFill('solid', fgColor='17365D')
        for col in sheet.columns:
            width = max(len(str(c.value or '')) for c in list(col)[:80])
            sheet.column_dimensions[get_column_letter(col[0].column)].width = min(44, max(16, width + 2))
        for row in sheet:
            for cell in row:
                if cell.data_type == 'f':
                    cell.data_type = 's'
    path = core.HERE / f'Nowgoal_{day}_多庄水位比分.xlsx'
    temp = path.with_suffix('.tmp.xlsx')
    wb.save(temp)
    try:
        temp.replace(path)
    except PermissionError:
        path = core.HERE / f'Nowgoal_{day}_多庄最新.xlsx'
        temp.replace(path)
    summary = {'file': str(path), 'source_status': source_status, **counts}
    (core.HERE / 'multi_summary.json').write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding='utf-8')
    print(json.dumps(summary, ensure_ascii=True), flush=True)
    return summary


def run(args):
    cycle_started = time.monotonic()
    core.OFFLINE = args.offline
    core.SOURCE_LIMITED = False
    path = core.HERE / f'multi_collection_{args.date}.json'
    seed = path if path.exists() else core.HERE / f'collection_{args.date}.json'
    matches = json.loads(seed.read_text(encoding='utf-8'))
    import_cached(matches)
    source_status = '已整理缓存；缺失的公司或赛前历史待补采'
    if not args.offline:
        try:
            payload = core.get_json(f'{core.BASE}/ajax/SoccerAjax?type=6&date={args.date}&order=time&timezone=8',
                                    f'fixtures_{args.date}.json', refresh=True)
            fixtures = core.parse_fixtures(payload)
            if any(not m['kickoff'].startswith(args.date + 'T') for m in fixtures):
                raise ValueError('Requested date does not match fixture dates')
            old = {m['id']: m for m in matches}
            for fresh in fixtures:
                saved = old.setdefault(fresh['id'], dict(fresh, collected_at=core.stamp(), companies=[]))
                saved.setdefault('comparison_at', saved['collected_at'])
                previous_state = saved['state']
                saved.update(fresh)
                saved['scores_at'] = core.stamp()
                if previous_state != fresh['state']:
                    saved['comparison_complete'] = False
                    saved['comparison_refresh_needed'] = True
                    # Refresh each book after match starts or ends; preserve previous records on failure.
                    saved['history_refresh_needed'] = [str(c['cid']) for c in saved.get('companies', []) if eligible(c)]
            matches = list(old.values())
            # Save and publish scores immediately; odds history must not delay score backfill.
            path.write_text(json.dumps(matches, ensure_ascii=False), encoding='utf-8')
            export(matches, args.date, '本轮比分已回填，多庄历史继续补采')
            odds_deadline = cycle_started + min(1200, max(0, args.interval - 120))
            budget_exhausted = False
            for index, match in enumerate(matches):
                if args.watch and time.monotonic() >= odds_deadline:
                    budget_exhausted = True
                    break
                mid = match['id']
                match.setdefault('comparison_at', match['collected_at'])
                if not match.get('comparison_complete') or match.get('comparison_refresh_needed'):
                    comp = core.get_json(core.odds_url(mid, match['state']), f'comp_{mid}.json', refresh=True)
                    match['companies'] = comp['Data']['mixodds']
                    match['comparison_complete'] = True
                    match['comparison_refresh_needed'] = False
                    match['comparison_at'] = core.stamp()
                histories = match.setdefault('multi_history', {})
                errors = match.setdefault('multi_errors', {})
                for company in match.get('companies', []):
                    if args.watch and time.monotonic() >= odds_deadline:
                        budget_exhausted = True
                        break
                    cid = str(company['cid'])
                    needs_refresh = cid in match.get('history_refresh_needed', [])
                    if match['state'] not in STARTED or not eligible(company) or (cid in histories and not needs_refresh):
                        continue
                    try:
                        hist = core.get_json(core.history_url(mid, cid), f'history_{mid}_{cid}.json', refresh=needs_refresh)
                        histories[cid] = hist['Data']
                        errors.pop(cid, None)
                        if needs_refresh:
                            match['history_refresh_needed'].remove(cid)
                    except Exception as error:
                        errors[cid] = str(error)
                        if core.SOURCE_LIMITED:
                            raise
                if index % 25 == 0:
                    path.write_text(json.dumps(matches, ensure_ascii=False), encoding='utf-8')
                    print(json.dumps({'processed_matches': index + 1, 'total': len(matches)}), flush=True)
            source_status = ('本轮比分已回填；剩余多庄历史留待下一轮补采' if budget_exhausted else
                             '本轮公开接口请求完成；无报价或无可核验历史仍保留为空')
        except Exception as error:
            source_status = '待补采：' + str(error)
            print(json.dumps({'request_stopped': source_status}, ensure_ascii=True), flush=True)
    path.write_text(json.dumps(matches, ensure_ascii=False), encoding='utf-8')
    return export(matches, args.date, source_status)


def next_cycle_delay(elapsed, interval):
    """Schedule from cycle start, skipping missed slots instead of overlapping workers."""
    interval = max(60, interval)
    return interval - elapsed % interval


def watch(args):
    status_path = core.HERE / 'multi_automation_status.json'
    def status(**fields):
        fields.update({'date': args.date, 'interval_seconds': args.interval, 'updated_at': core.stamp()})
        status_path.write_text(json.dumps(fields, ensure_ascii=False, indent=2), encoding='utf-8')
    if args.initial_delay:
        status(state='waiting', next_run=(datetime.now(core.TZ) + timedelta(seconds=args.initial_delay)).isoformat())
        time.sleep(max(0, args.initial_delay))
    while True:
        start = time.monotonic()
        status(state='running', last_started=core.stamp(),
               next_run=(datetime.now(core.TZ) + timedelta(seconds=max(60, args.interval))).isoformat())
        try:
            result = run(args)
            done = (result['pending_matches'] == 0 and result['comparison_pending_matches'] == 0
                    and result['history_pending_rows'] == 0)
            if not args.watch or done:
                status(state='complete', last_finished=core.stamp())
                break
        except Exception as error:
            print(json.dumps({'cycle_error': str(error), 'time': core.stamp()}, ensure_ascii=True), flush=True)
            if not args.watch:
                raise
        delay = next_cycle_delay(time.monotonic() - start, args.interval)
        next_run = (datetime.now(core.TZ) + timedelta(seconds=delay)).isoformat()
        status(state='waiting', last_finished=core.stamp(), next_run=next_run)
        print(json.dumps({'next_run': next_run, 'interval_seconds': args.interval}), flush=True)
        time.sleep(delay)


if __name__ == '__main__':
    parser = argparse.ArgumentParser()
    parser.add_argument('--date', default=datetime.now(core.TZ).date().isoformat())
    parser.add_argument('--offline', action='store_true')
    parser.add_argument('--watch', action='store_true')
    parser.add_argument('--interval', type=int, default=1800)
    parser.add_argument('--initial-delay', type=int, default=0)
    args = parser.parse_args()
    watch(args)
