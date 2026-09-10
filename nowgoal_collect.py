"""Collect public fixture/odds data; export Excel and optionally backfill results."""
import argparse
import ast
from concurrent.futures import ThreadPoolExecutor, as_completed
from datetime import datetime, timedelta, timezone
import json
from pathlib import Path
import re
import time
from collections import Counter

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from fetch_nowgoal import fetch, ROOT

BASE = 'https://live11.nowgoal26.com'
HERE = Path(__file__).resolve().parent
TZ = timezone(timedelta(hours=8))
SOURCE_LIMITED = False
OFFLINE = False
STATES = {-1: '完场', -10: '取消', -11: '待定', -12: '腰斩', -13: '中断', -14: '推迟', 0: '未开赛', 1: '上半场', 2: '中场', 3: '下半场', 4: '加时', 5: '点球'}

def stamp():
    return datetime.now(TZ).isoformat(timespec='seconds')

def literal_array(raw):
    # Parse only scalar literals. Never execute the downloaded JavaScript.
    parts, start, quote, escape = [], 0, None, False
    for i, char in enumerate(raw):
        if escape:
            escape = False
        elif quote and char == '\\':
            escape = True
        elif quote and char == quote:
            quote = None
        elif not quote and char in "\"'":
            quote = char
        elif not quote and char == ',':
            parts.append(raw[start:i]); start = i + 1
    parts.append(raw[start:])
    values = []
    for part in parts:
        part = part.strip()
        if part in ('', 'null', 'undefined'):
            values.append(None)
        elif part in ('true', 'false'):
            values.append(part == 'true')
        else:
            values.append(ast.literal_eval(part))
    return values

def parse_fixtures(payload):
    if payload.get('ErrCode') != 0:
        raise ValueError('Fixture API error: ' + str(payload.get('ErrCode')))
    source = payload['Data']
    tables = {'A': {}, 'B': {}}
    for table, idx, raw in re.findall(r'^([AB])\[(\d+)\]=\[(.*?)\];', source, re.M):
        tables[table][int(idx)] = literal_array(raw)
    expected = int(re.search(r'var matchcount=(\d+)', source)[1])
    if len(tables['A']) != expected:
        raise ValueError('Fixture count mismatch')
    rows = []
    for a in tables['A'].values():
        date = [int(v) for v in a[6].split(',')]
        date[1] += 1
        kick = datetime(*date, tzinfo=timezone.utc).astimezone(TZ)
        b = tables['B'][a[1]]
        rows.append({'id': a[0], 'league': b[2], 'home': a[4], 'away': a[5],
                     'kickoff': kick.isoformat(), 'state': a[7], 'home_score': a[8],
                     'away_score': a[9], 'half_home': a[10], 'half_away': a[11],
                     'explain': a[19:24]})
    if len({r['id'] for r in rows}) != expected:
        raise ValueError('Duplicate fixture IDs')
    return rows

def get_json(url, filename, refresh=True):
    global SOURCE_LIMITED
    path = ROOT / filename
    if (not refresh or OFFLINE) and path.exists():
        data = json.loads(path.read_text(encoding='utf-8'))
        if data.get('ErrCode') == 0:
            return data
    if SOURCE_LIMITED or OFFLINE:
        raise ValueError('待补采：源站限制请求或本次仅整理已下载数据')
    for attempt in range(3):
        try:
            time.sleep(1.5)
            data = json.loads(fetch(url, filename))
            if data.get('code') in (1001, 100401):
                SOURCE_LIMITED = True
                print(json.dumps({'source_limit': data['code'], 'action': 'stop_network_until_next_run'}), flush=True)
                raise ValueError('Source temporarily restricts this request: ' + str(data))
            if data.get('ErrCode') != 0:
                raise ValueError('API error: ' + str(data))
            return data
        except Exception:
            if attempt == 2 or SOURCE_LIMITED:
                raise
            time.sleep(2 * (attempt + 1))

def odds_url(mid, state):
    return f'{BASE}/ajax/soccerajax?type=14&t=1&id={mid}&h=0&s={state}'

def history_url(mid, cid):
    return f'{BASE}/ajax/soccerajax?type=14&id={mid}&t=20&cid={cid}&h=0&r1=0&r2=0&r3=0'

def complete(odds):
    return bool(odds) and all(odds.get(k) not in ('', None) for k in ('u', 'g', 'd'))

def collect_match(match, cid, refresh):
    out = dict(match)
    out['collected_at'] = stamp()
    try:
        comp = get_json(odds_url(match['id'], match['state']), f"comp_{match['id']}.json", refresh)
        out['companies'] = comp['Data']['mixodds']
        selected = next((c for c in out['companies'] if c['cid'] == cid), {})
        out['selected'] = selected
        if complete(selected.get('ah', {}).get('f')) or complete(selected.get('ah', {}).get('l')):
            hist = get_json(history_url(match['id'], cid), f"history_{match['id']}_{cid}.json", refresh)
            out['history'] = hist['Data']
    except Exception as error:
        out['error'] = str(error)
        # Already downloaded public bulk data remains useful when detail requests fail.
        bulk = ROOT / f'odds{cid}.txt'
        if not out.get('selected') and bulk.exists():
            line = next((s for s in bulk.read_text(encoding='utf-8').split('$')
                         if s.startswith(str(match['id']) + '!')), None)
            if line:
                parts = line.split('!')
                company = {'cid': cid}
                for index, market in enumerate(('ah', 'euro', 'ou'), 1):
                    values = parts[index].split(',')
                    company[market] = {key: dict(zip(('u', 'g', 'd'), values[start:start + 3]))
                                       for key, start in [('f', 0), ('l', 3), ('r', 6)]}
                out['selected'] = company
                out['companies'] = [company]
                out['bulk_snapshot'] = True
                out['collected_at'] = datetime.fromtimestamp(bulk.stat().st_mtime, TZ).isoformat()
    time.sleep(0.12)
    return out

def num(value):
    return float(value) if value not in ('', None) else None

def trio(odds):
    return [num(odds.get(k)) for k in ('u', 'g', 'd')] if odds else [None] * 3

def prematch(match, market='ah'):
    kickoff = datetime.fromisoformat(match['kickoff']).timestamp()
    records = match.get('history', {}).get(market, [])
    return sorted((r for r in records if r.get('type') in (1, 2)
                   and not r.get('close') and complete(r.get('odds'))
                   and r.get('mt', kickoff) < kickoff), key=lambda r: r['mt'])

def report(matches, day, cid):
    wb = Workbook()
    ws = wb.active
    ws.title = '赛事水位比分'
    headers = ['比赛ID', '联赛', '开赛时间(北京时间)', '主队', '客队', '状态', '赔率公司',
               '初盘主水', '初盘盘口(正=主让)', '初盘客水', '初盘最早匹配时间(北京时间)',
               '赛前最后主水', '赛前最后盘口(正=主让)', '赛前最后客水', '赛前最后记录时间(北京时间)',
               '当前赛前主水', '当前赛前盘口(正=主让)', '当前赛前客水', '水位核验',
               '终场主队比分', '终场客队比分', '终场比分', '半场比分', '比分回填状态',
               '当前比分', '采集时间', '来源', '备注']
    ws.append(headers)
    company_names = {8: 'Bet365', 3: 'Crown', 31: 'Sbobet', 50: '1xBet', 17: 'M88',
                     24: '12Bet', 42: '18Bet', 12: 'Easybet', 1: 'Macauslot',
                     4: 'Ladbrokes', 14: 'Vcbet', 19: 'Interwetten'}
    counters = Counter()
    history_rows, all_rows = [], []
    for match in sorted(matches, key=lambda m: (m['kickoff'], m['id'])):
        mid, state = match['id'], match['state']
        selected = match.get('selected', {}).get('ah', {})
        initial = selected.get('f', {})
        live = selected.get('l', {})
        pre = prematch(match)
        closing, opening_time, closing_time = {}, None, None
        if pre:
            same = [r for r in pre if trio(r['odds']) == trio(initial)]
            if same:
                opening_time = datetime.fromtimestamp(same[0]['mt'], TZ).isoformat()
        started = state in (-1, 1, 2, 3, 4, 5)
        note = []
        if started and pre:
            closing = pre[-1]['odds']
            closing_time = datetime.fromtimestamp(pre[-1]['mt'], TZ).isoformat()
            verified = '已核对赛前历史时间'
            if complete(live) and trio(live) != trio(closing):
                note.append('历史最后赛前记录与页面即时盘不同；临场列采用历史记录')
        elif started:
            verified = '缺少可核验赛前历史'
            if complete(live):
                note.append('页面提供即时盘，但缺少赛前历史时间；保留在当前赛前列，临场列留空')
        elif state == 0:
            verified = '未开赛，待赛前最后记录'
        else:
            verified = '异常赛况，待确认'
        if not complete(initial):
            note.append('该公司未提供完整亚洲让球初盘')
        if match.get('error'):
            note.append('请求失败：' + match['error'])
        if match.get('bulk_snapshot'):
            note.append('水位来自已下载的页面汇总快照，单场历史仍待补采')
        final = state == -1 and match['home_score'] is not None and match['away_score'] is not None
        current_score = f"{match['home_score']}-{match['away_score']}" if started else None
        final_score = current_score if final else None
        half = f"{match['half_home']}-{match['half_away']}" if started and match['half_home'] is not None and match['half_away'] is not None else None
        backfill = '已回填' if final else ('待回填' if state >= 0 else '异常赛况，无终场比分')
        row = [mid, match['league'], match['kickoff'], match['home'], match['away'], STATES.get(state, str(state)),
               company_names.get(cid, str(cid)), *trio(initial), opening_time, *trio(closing), closing_time,
               *trio(live), verified, match['home_score'] if final else None, match['away_score'] if final else None,
               final_score, half, backfill, current_score, match['collected_at'], f'{BASE}/oddscomp/{mid}', '; '.join(note)]
        ws.append(row)
        counters['matches'] += 1
        counters['finished'] += final
        counters['initial_ah'] += complete(initial)
        counters['verified_closing_ah'] += complete(closing)
        counters['pending'] += state >= 0
        counters['errors'] += bool(match.get('error'))
        counters['bulk_snapshot'] += bool(match.get('bulk_snapshot'))
        for record in pre:
            history_rows.append([mid, company_names.get(cid, cid), datetime.fromtimestamp(record['mt'], TZ).isoformat(),
                                 *trio(record['odds']), record['type'], f'{BASE}/oddscomp/{mid}'])
        for company in match.get('companies', []):
            for market, title in [('ah', '亚洲让球(香港水位)'), ('ou', '大小球(香港水位)'), ('euro', '胜平负(欧洲赔率)')]:
                odds = company.get(market, {})
                if complete(odds.get('f')) or complete(odds.get('l')):
                    all_rows.append([mid, match['league'], match['home'], match['away'], company_names.get(company['cid'], f"公司ID {company['cid']}"),
                                     title, *trio(odds.get('f')), *trio(odds.get('l')), final_score, match['collected_at']])
    for title, columns, rows in [
        ('赛前历史核验', ['比赛ID', '赔率公司', '记录时间(北京时间)', '主水', '盘口(正=主让)', '客水', '网站赛前类型', '来源'], history_rows),
        ('各公司盘口快照', ['比赛ID', '联赛', '主队', '客队', '赔率公司', '盘口类型', '初盘主胜或大球水位', '初盘盘口或平赔', '初盘客胜或小球水位', '页面即时主胜或大球水位', '页面即时盘口或平赔', '页面即时客胜或小球水位', '终场比分', '采集时间'], all_rows)]:
        sheet = wb.create_sheet(title); sheet.append(columns)
        for row in rows:
            sheet.append(row)
    notes = wb.create_sheet('说明')
    for row in [
        ['项目', '说明'], ['范围', f'北京时间 {day} 00:00:00 至次日 00:00:00 的网站赛程列表全部赛事；含无赔率赛事。'],
        ['来源', BASE + '/football/fixture/'], ['采集方式', '公开赛程API + 单场赔率比较API + 默认公司赛前历史。原始响应保存在evidence目录。'],
        ['初盘', '网站赔率比较数据的 f 字段。初盘时间仅在历史中找到相同水位时填写，表示最早可见匹配时间。'],
        ['赛前最后水位', '仅对已开赛/完场赛事，选择网站标记为赛前(type=1或2)、未封盘、时间严格早于计划开赛的最后记录；无可核验历史则留空。'],
        ['时间限制', '历史核验按网站计划开赛时间判断；网站未提供实际开球时间，延迟开球可能导致所取记录不是实际开球前最后一条。'],
        ['即时盘', '网站 l 字段，与滚球 r 字段分开；未开赛时仍会变化，不能视为最终临场水位。'],
        ['水位', '亚洲让球及大小球使用香港水位；胜平负使用欧洲十进制赔率。'],
        ['让球符号', '网站数值正数表示主队让球，负数表示主队受让；例如-1.5表示主队受让一球半。'],
        ['比分', '仅state=-1(完场)填写终场比分；其他赛事保留当前比分及待回填状态。比分使用网站全场比分字段，不擅自合并点球结果。'],
        ['缺失值', '空白表示源站未提供/尚未发生/无法核验，绝不补零。'],
        ['各公司快照', '附表保留接口返回的其他公司与盘口，但其即时盘未逐公司核验历史，不应一律当作临场盘。'],
        ['自动更新', 'python nowgoal_collect.py --date ' + day + ' --watch；每5分钟检查赛果，开赛及完场时重新获取水位。持续至无待完成赛事或72小时上限。'],
        ['采集统计', json.dumps(dict(counters), ensure_ascii=False)], ['更新时间', stamp()]]:
        notes.append(row)
    for sheet in wb:
        sheet.freeze_panes = 'A2'
        sheet.auto_filter.ref = sheet.dimensions
        sheet.row_dimensions[1].height = 32
        for cell in sheet[1]:
            cell.font = Font(bold=True, color='FFFFFF')
            cell.fill = PatternFill('solid', fgColor='17365D')
        for column in sheet.columns:
            index = column[0].column
            width = max(len(str(c.value or '')) for c in list(column)[:100])
            sheet.column_dimensions[get_column_letter(index)].width = min(48, max(15, width + 2))
        for row in sheet:
            for cell in row:
                if cell.data_type == 'f':
                    cell.data_type = 's'  # Source text must never become an Excel formula.
    path = HERE / f'Nowgoal_{day}_赛事水位比分.xlsx'
    temp = path.with_suffix('.tmp.xlsx')
    wb.save(temp)
    try:
        temp.replace(path)
    except PermissionError:
        path = HERE / f'Nowgoal_{day}_最新.xlsx'
        temp.replace(path)
    (HERE / 'collection_summary.json').write_text(json.dumps(dict(counters), indent=2), encoding='utf-8')
    print(json.dumps({'report': str(path), **counters}, ensure_ascii=False), flush=True)
    return counters

def run(args):
    global OFFLINE, SOURCE_LIMITED
    OFFLINE = args.offline
    SOURCE_LIMITED = False
    state_path = HERE / f'collection_{args.date}.json'
    previous = json.loads(state_path.read_text(encoding='utf-8')) if state_path.exists() else []
    old = {m['id']: m for m in previous}
    url = f'{BASE}/ajax/SoccerAjax?type=6&date={args.date}&order=time&timezone=8'
    fixture_file = f'fixtures_{args.date}.json'
    if args.offline:
        candidates = [ROOT / fixture_file, ROOT / 'today.json']
        payload = next((json.loads(p.read_text(encoding='utf-8')) for p in candidates
                        if p.exists() and json.loads(p.read_text(encoding='utf-8')).get('ErrCode') == 0), None)
        if payload is None:
            raise ValueError('No successful fixture snapshot available')
        fixtures = parse_fixtures(payload)
    else:
        fixtures = parse_fixtures(get_json(url, fixture_file, refresh=True))
    if any(not m['kickoff'].startswith(args.date + 'T') for m in fixtures):
        raise ValueError('Fixture snapshot does not belong to the requested Beijing date')
    # Keep the requested date fixed, even when a postponed match moves out of it.
    if args.limit:
        fixtures = fixtures[:args.limit]
    matches, jobs = [], []
    for match in fixtures:
        saved = old.get(match['id'])
        if args.watch and saved and not saved.get('error') and saved['state'] == match['state']:
            merged = dict(saved); merged.update(match); merged['collected_at'] = stamp()
            matches.append(merged)
        else:
            jobs.append(match)
    with ThreadPoolExecutor(max_workers=1) as pool:
        futures = {pool.submit(collect_match, m, args.company, args.refresh or args.watch): m for m in jobs}
        for i, future in enumerate(as_completed(futures), 1):
            updated = future.result()
            saved = old.get(updated['id'])
            if updated.get('error') and saved:
                for key in ('history', 'selected', 'companies'):
                    if not updated.get(key) and saved.get(key):
                        updated[key] = saved[key]
                updated['collected_at'] = saved['collected_at']
            matches.append(updated)
            if i % 30 == 0:
                print(json.dumps({'progress': i, 'total': len(jobs)}), flush=True)
    current_ids = {m['id'] for m in matches}
    matches.extend(m for mid, m in old.items() if mid not in current_ids)
    state_path.write_text(json.dumps(matches, ensure_ascii=False, indent=2), encoding='utf-8')
    return report(matches, args.date, args.company)

if __name__ == '__main__':
    parser = argparse.ArgumentParser()
    parser.add_argument('--date', default=datetime.now(TZ).date().isoformat())
    parser.add_argument('--company', type=int, default=8)
    parser.add_argument('--refresh', action='store_true')
    parser.add_argument('--watch', action='store_true')
    parser.add_argument('--offline', action='store_true')
    parser.add_argument('--interval', type=int, default=300)
    parser.add_argument('--limit', type=int, default=0)
    args = parser.parse_args()
    deadline = time.monotonic() + 72 * 3600
    while True:
        try:
            result = run(args)
            if not args.watch or (result['pending'] == 0 and result['errors'] == 0):
                break
        except Exception as error:
            print(json.dumps({'fatal_error': str(error), 'time': stamp()}, ensure_ascii=False), flush=True)
            if not args.watch:
                raise
        if time.monotonic() >= deadline:
            break
        time.sleep(max(60, args.interval))
