"""Does replacing the multiplicative de-vig with Shin's method change backtest accuracy?
Two changes tested together (both additive-only, v8_backtest_pipeline.py untouched):
  1. water_sig: Shin-devigged over/under probability instead of multiplicative.
  2. euro_sig: a REAL devigged draw-probability drift (Shin, using home+draw+away odds),
     replacing the current crude "raw draw-odds move" proxy that isn't devigged at all.
Tested on the 8-day/~3939-match dataset with a genuine date-based train/test split
(first half of calendar dates = train, second half = test), matching how the
"guess-under + draw-prob-rising" cross-signal was validated earlier this session.
"""
import statistics as st
import sys
from collections import defaultdict, Counter

import openpyxl
import shin

sys.path.insert(0, r"C:\Users\Administrator\Desktop\zpp")
import importlib.util
spec = importlib.util.spec_from_file_location("pipeline", r"C:\Users\Administrator\Desktop\zpp\v8_backtest_pipeline.py")
pipeline = importlib.util.module_from_spec(spec)
spec.loader.exec_module(pipeline)

XLSX = r"C:\Users\Administrator\Desktop\zpp\Nowgoal_2026-09-02_至_2026-09-09_多庄水位比分.xlsx"
matches, ah_by_match, ou_by_match, eu_by_match_orig = pipeline.load_workbook_data(XLSX, min_companies=2)

wb = openpyxl.load_workbook(XLSX, data_only=True)

# Recover kickoff date (dropped by load_workbook_data) for the train/test split.
kickoff_by_mid = {}
for r in wb['全部赛事与比分'].iter_rows(min_row=2, values_only=True):
    mid, league, kickoff = r[0], r[1], r[2]
    if mid in matches and kickoff:
        kickoff_by_mid[mid] = str(kickoff)[:10]

# Recover the euro home/away odds that load_workbook_data() drops (only keeps draw).
eu_full = defaultdict(list)
for r in wb['各庄大小球与欧赔'].iter_rows(min_row=2, values_only=True):
    (mid, league, home, away, cid, cname, ptype,
     open_a, open_mid, open_b, live_a, live_mid, live_b, fscore) = r
    if mid not in matches or ptype != '胜平负(欧洲赔率)':
        continue
    if None in (open_a, open_mid, open_b, live_a, live_mid, live_b):
        continue
    try:
        vals = [float(open_a), float(open_mid), float(open_b), float(live_a), float(live_mid), float(live_b)]
    except (TypeError, ValueError):
        continue
    if any(v <= 0 for v in vals):
        continue
    eu_full[mid].append(dict(open_home=vals[0], open_draw=vals[1], open_away=vals[2],
                              live_home=vals[3], live_draw=vals[4], live_away=vals[5]))


def shin_pair(over_water, under_water):
    probs = shin.calculate_implied_probabilities([1 + over_water, 1 + under_water])
    return probs[0], probs[1]


def shin_draw_prob(home_odds, draw_odds, away_odds):
    return shin.calculate_implied_probabilities([home_odds, draw_odds, away_odds])[1]


def compute_shin_variant(mid):
    ou_rows = ou_by_match.get(mid, [])
    ah_rows = ah_by_match.get(mid, [])
    if len(ou_rows) < 2:
        return None
    line_signals, water_signals = [], []
    for row in ou_rows:
        line_signals.append(pipeline.sign(row['live_line'] - row['open_line'], 0.001))
        try:
            p_over_open, _ = shin_pair(row['open_over'], row['open_under'])
            p_over_live, _ = shin_pair(row['live_over'], row['live_under'])
        except Exception:
            continue
        water_signals.append(pipeline.sign(p_over_live - p_over_open, 0.01))
    if not water_signals:
        return None
    line_sig = st.mean(line_signals)
    water_sig = st.mean(water_signals)

    euro_sig = 0.0
    eu_rows_full = eu_full.get(mid, [])
    if eu_rows_full:
        euro_deltas = []
        for r in eu_rows_full:
            try:
                p_draw_open = shin_draw_prob(r['open_home'], r['open_draw'], r['open_away'])
                p_draw_live = shin_draw_prob(r['live_home'], r['live_draw'], r['live_away'])
            except Exception:
                continue
            euro_deltas.append(pipeline.sign(p_draw_live - p_draw_open, 0.01))
        if euro_deltas:
            euro_sig = st.mean(euro_deltas)

    handicap_sig = 0.0
    if ah_rows:
        deltas = [row['last_line'] - row['open_line'] for row in ah_rows]
        handicap_sig = st.mean([pipeline.sign(d, 0.001) for d in deltas])

    weights = {'line': 20, 'water': 15, 'euro': 10, 'handicap': 10}
    total_w = sum(weights.values())
    sig = {'line': line_sig, 'water': water_sig, 'euro': euro_sig, 'handicap': handicap_sig}
    composite = sum(sig[k] * w for k, w in weights.items()) / total_w
    dirs = [d for d in (1 if x > 0 else (-1 if x < 0 else 0) for x in line_signals) if d != 0]
    consistency_dir = Counter(dirs).most_common(1)[0][0] if dirs else 0
    composite += (0.05 if consistency_dir > 0 else (-0.05 if consistency_dir < 0 else 0)) * (5 / total_w) * 20
    composite = max(-1, min(1, composite))
    prob_over = max(20, min(80, 50 + composite * 20))
    direction = 'over' if prob_over >= 50 else 'under'
    return dict(direction=direction, euro_sig=euro_sig, has_full_euro=bool(eu_rows_full))


rows = []
for mid, m in matches.items():
    if m['status'] != '完场' or m['total_goals'] is None or mid not in kickoff_by_mid:
        continue
    orig = pipeline.compute_odds_drift_signal(mid, ou_by_match.get(mid, []), ah_by_match.get(mid, []),
                                               eu_by_match_orig.get(mid, []))
    shin_r = compute_shin_variant(mid)
    if orig is None or shin_r is None:
        continue
    line = orig['consensus_line']
    if m['total_goals'] == line:
        continue  # push, excluded from hit-rate denominator
    actual = 'over' if m['total_goals'] > line else 'under'
    rows.append(dict(mid=mid, date=kickoff_by_mid[mid], line=line, actual=actual,
                      orig_dir=orig['direction'], shin_dir=shin_r['direction'],
                      has_full_euro=shin_r['has_full_euro']))

dates = sorted(set(r['date'] for r in rows))
mid_date = dates[len(dates) // 2]
train = [r for r in rows if r['date'] < mid_date]
test = [r for r in rows if r['date'] >= mid_date]
print(f"日期范围: {dates[0]} ~ {dates[-1]}, 切分点: {mid_date}")
print(f"总可评分场次: {len(rows)}  train={len(train)}  test={len(test)}")
print(f"其中有完整欧赔三项数据、euro_sig可用Shin重算的: {sum(1 for r in rows if r['has_full_euro'])}\n")


def report(label, grp):
    if not grp:
        print(f"{label}: 无样本"); return
    orig_hit = sum(1 for r in grp if r['orig_dir'] == r['actual'])
    shin_hit = sum(1 for r in grp if r['shin_dir'] == r['actual'])
    agree = sum(1 for r in grp if r['orig_dir'] == r['shin_dir'])
    print(f"{label} (n={len(grp)}): 原方法命中率={orig_hit/len(grp)*100:.2f}%  "
          f"Shin方法命中率={shin_hit/len(grp)*100:.2f}%  两方法方向一致={agree/len(grp)*100:.1f}%")


report('train', train)
report('test', test)
report('全部', rows)

# Where the two methods disagree, which one was right more often? (only informative subset)
disagree = [r for r in rows if r['orig_dir'] != r['shin_dir']]
if disagree:
    orig_right = sum(1 for r in disagree if r['orig_dir'] == r['actual'])
    print(f"\n两方法方向不一致的场次(n={len(disagree)}): 原方法对{orig_right}次, "
          f"Shin方法对{len(disagree)-orig_right}次")
