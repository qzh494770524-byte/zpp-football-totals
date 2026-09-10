"""Reproducible descriptive over/under analysis of the local Nowgoal snapshot."""
import argparse
from collections import Counter, defaultdict
from datetime import datetime
import html
import json
import math
from pathlib import Path
import re
import statistics
import time

import numpy as np
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

import nowgoal_collect as core
import nowgoal_multi as multi

RULES = 'https://help.bet365.es/s/en-es/sportsrules/soccer/goal-line'
TIME_RULES = 'https://help.bet365.com/s/en/sportsrules/soccer/result-event-half-time'
LABELS = {-2: '全输', -1: '半输', 0: '走盘', 1: '半赢', 2: '全赢'}


def read_json(path):
    for attempt in range(3):
        try:
            return json.loads(path.read_text(encoding='utf-8'))
        except (json.JSONDecodeError, PermissionError):
            if attempt == 2:
                raise
            time.sleep(.2)


def quote(value):
    try:
        over, line, under = [float(value[k]) for k in ('u', 'g', 'd')]
        if not all(math.isfinite(v) for v in (over, line, under)) or over <= 0 or under <= 0 or line < 0:
            return None
        if abs(line * 4 - round(line * 4)) > 1e-8:
            return None
        return {'over': over, 'line': line, 'under': under}
    except (KeyError, TypeError, ValueError):
        return None


def settle(total, line, water, side='over'):
    if side not in ('over', 'under') or total < 0 or int(total) != total or water <= 0:
        raise ValueError('Invalid settlement input')
    if line < 0 or abs(line * 4 - round(line * 4)) > 1e-8:
        raise ValueError('Only whole, half and quarter goal lines are supported')
    lower, upper = math.floor(line * 2) / 2, math.ceil(line * 2) / 2
    parts = [lower, upper]
    signs = [((total > p) - (total < p)) * (1 if side == 'over' else -1) for p in parts]
    profit = sum(water if sign > 0 else -1 if sign < 0 else 0 for sign in signs) / 2
    return {'grade': sum(signs), 'result': LABELS[sum(signs)], 'profit': profit}


def exclusion(match):
    if re.search(r'futsal|beach|esoccer|e-soccer|simulat|indoor|kings league|baller league|\b\d+ mins', match['league'], re.I):
        return '非标准赛制/模拟赛事'
    note = ' '.join(str(v) for v in match.get('explain', []))
    if re.search(r'120\s*min|Pen\[|extra.?time|overtime|\bAET\b|abandon|award|walkover', note, re.I):
        return '加时/点球或特殊赛果备注，常规时间比分待核实'
    return ''


def category(match):
    text = ' '.join(match[k] for k in ('league', 'home', 'away'))
    if re.search(r'\(W\)|women|bayanlar|feminin|femenin|frauen|ladies|damer', text, re.I):
        return '女子赛事（含青年）'
    if re.search(r'\bU[- ]?(?:1\d|2[0-3])\b|youth|junior|primavera', text, re.I):
        return '青年赛事（不含女子）'
    return '其他赛事（含成年与预备队）'


def freeze_snapshot(out):
    cutoff = time.time()
    matches = {}
    for path in [core.HERE / 'history_collection_2026-09-02_2026-09-08.json', core.HERE / 'multi_collection_2026-09-09.json']:
        for match in read_json(path):
            matches[match['id']] = match
    games, quotes, errors = [], [], []
    for match in matches.values():
        if match['state'] != -1 or match.get('home_score') is None or match.get('away_score') is None:
            continue
        if any(not isinstance(match[k], int) or match[k] < 0 for k in ('home_score', 'away_score')):
            errors.append({'id': match['id'], 'error': '比分无效'})
            continue
        mid = match['id']
        game = {k: match[k] for k in ('id', 'league', 'home', 'away', 'kickoff', 'home_score', 'away_score')}
        game.update(total=match['home_score'] + match['away_score'], date=match['kickoff'][:10],
                    category=category(match), exclude=exclusion(match), note=match.get('explain', []))
        games.append(game)
        companies = match.get('companies', [])
        comp_file = core.ROOT / f'comp_{mid}.json'
        if comp_file.exists() and comp_file.stat().st_mtime <= cutoff:
            try:
                comp = read_json(comp_file)
                if comp.get('ErrCode') == 0:
                    companies = comp['Data']['mixodds']
            except (ValueError, KeyError) as error:
                errors.append({'id': mid, 'error': str(error)})
        for company in companies:
            cid = company['cid']
            initial = quote(company.get('ou', {}).get('f'))
            latest = quote(company.get('ou', {}).get('l'))
            closing, closing_at = None, None
            hist_file = core.ROOT / f'history_{mid}_{cid}.json'
            if hist_file.exists() and hist_file.stat().st_mtime <= cutoff:
                try:
                    hist = read_json(hist_file)
                    if hist.get('ErrCode') == 0:
                        records = [r for r in hist.get('Data', {}).get('ou', [])
                                   if r.get('type') in (1, 2) and not r.get('close') and quote(r.get('odds'))
                                   and r.get('mt', float('inf')) < datetime.fromisoformat(match['kickoff']).timestamp()]
                        if records:
                            last = max(records, key=lambda r: r['mt'])
                            closing, closing_at = quote(last['odds']), datetime.fromtimestamp(last['mt'], core.TZ).isoformat()
                except (ValueError, KeyError) as error:
                    errors.append({'id': mid, 'cid': cid, 'error': str(error)})
            if initial or closing:
                quotes.append({'id': mid, 'cid': cid, 'company': multi.name(cid), 'initial': initial,
                               'page_latest': latest, 'closing': closing, 'closing_at': closing_at,
                               'comparison_file': str(comp_file), 'history_file': str(hist_file) if closing else None})
    result = {'captured_at': datetime.fromtimestamp(cutoff, core.TZ).isoformat(), 'games': games, 'quotes': quotes, 'read_errors': errors}
    (out / 'input_snapshot.json').write_text(json.dumps(result, ensure_ascii=False), encoding='utf-8')
    return result


def wilson(successes, n):
    if not n:
        return (None, None)
    z = 1.95996398454
    p = successes / n
    center = (p + z*z/(2*n)) / (1 + z*z/n)
    margin = z * math.sqrt(p*(1-p)/n + z*z/(4*n*n)) / (1 + z*z/n)
    return center - margin, center + margin


def bootstrap(profits):
    if not profits:
        return (None, None)
    rng = np.random.default_rng(20260909)
    a = np.array(profits)
    sample = np.concatenate([a[rng.integers(0, len(a), size=(250, len(a)))].mean(axis=1) for _ in range(12)])
    return tuple(float(v) for v in np.quantile(sample, [.025, .975]))


def game_stats(games):
    n = len(games)
    totals = [g['total'] for g in games]
    return {'n': n, 'mean_goals': statistics.mean(totals) if n else None,
            'median_goals': statistics.median(totals) if n else None,
            'over15': sum(v >= 2 for v in totals)/n if n else None,
            'over25': sum(v >= 3 for v in totals)/n if n else None,
            'over35': sum(v >= 4 for v in totals)/n if n else None,
            'zero': sum(v == 0 for v in totals)/n if n else None,
            'btts': sum(g['home_score'] > 0 and g['away_score'] > 0 for g in games)/n if n else None,
            'over25_ci': wilson(sum(v >= 3 for v in totals), n)}


def market_stats(rows, phase='initial', confidence=False):
    outcomes = {side: [settle(r['total'], r[phase]['line'], r[phase][side], side) for r in rows] for side in ('over', 'under')}
    result = {'n': len(rows)}
    for side, values in outcomes.items():
        counts = Counter(v['result'] for v in values)
        profits = [v['profit'] for v in values]
        result[side] = {'counts': dict(counts), 'roi': statistics.mean(profits) if profits else None,
                        'net_units': sum(profits), 'roi_ci': bootstrap(profits) if confidence else None}
    return result


def analyze(snapshot):
    games = [g for g in snapshot['games'] if not g['exclude']]
    game_map = {g['id']: g for g in games}
    rows = [dict(game_map[q['id']], **q) for q in snapshot['quotes'] if q['id'] in game_map]
    initial = [r for r in rows if r['initial']]
    closing = [r for r in rows if r['closing']]
    initial_ids = {r['id'] for r in initial}
    closing_ids = {r['id'] for r in closing}
    by_company = []
    for cid in sorted({r['cid'] for r in initial}):
        book_rows = [r for r in initial if r['cid'] == cid]
        by_company.append({'cid': cid, 'company': multi.name(cid), **market_stats(book_rows, confidence=True)})
    daily = []
    for day in sorted({g['date'] for g in games}):
        group = [g for g in games if g['date'] == day]
        daily.append({'date': day, **game_stats(group), 'initial_games': sum(g['id'] in initial_ids for g in group),
                      'closing_games': sum(g['id'] in closing_ids for g in group)})
    leagues = []
    for league in sorted({g['league'] for g in games}):
        group = [g for g in games if g['league'] == league]
        leagues.append({'league': league, **game_stats(group), 'initial_games': sum(g['id'] in initial_ids for g in group)})
    categories = [{'category': c, **game_stats([g for g in games if g['category'] == c])} for c in sorted({g['category'] for g in games})]
    reference = [r for r in initial if r['cid'] == 8]
    lines = [{'line': line, **market_stats([r for r in reference if r['initial']['line'] == line])}
             for line in sorted({r['initial']['line'] for r in reference})]
    movements = []
    for cid in (8, 3, 31, 50):
        paired = [r for r in rows if r['cid'] == cid and r['initial'] and r['closing']]
        for label, sign in [('升盘', 1), ('不变', 0), ('降盘', -1)]:
            selected = [r for r in paired if (r['closing']['line'] > r['initial']['line']) - (r['closing']['line'] < r['initial']['line']) == sign]
            movements.append({'company': multi.name(cid), 'move': label, 'initial': market_stats(selected),
                              'closing': market_stats(selected, 'closing')})
    common_ids = set.intersection(*[{r['id'] for r in initial if r['cid'] == cid} for cid in (8, 3, 31, 50)])
    common = [{'company': multi.name(cid), **market_stats([r for r in initial if r['cid'] == cid and r['id'] in common_ids], confidence=True)}
              for cid in (8, 3, 31, 50)]
    summary = {'captured_at': snapshot['captured_at'], 'finished_raw': len(snapshot['games']),
               'excluded': len(snapshot['games'])-len(games), 'overall': game_stats(games),
               'initial_games': len(initial_ids), 'initial_quotes': len(initial), 'closing_games': len(closing_ids),
               'closing_quotes': len(closing), 'company_count': len(by_company),
               'quoted_games_stats': game_stats([g for g in games if g['id'] in initial_ids]),
               'unquoted_games_stats': game_stats([g for g in games if g['id'] not in initial_ids]),
               'daily': daily, 'categories': categories, 'leagues': leagues, 'companies': by_company,
               'reference_initial_lines': lines, 'movements': movements, 'common_four_company': common}
    return summary, games, rows


def excel(out, summary, snapshot, games, rows):
    wb = Workbook(); wb.remove(wb.active)
    def sheet(title, headers, data):
        s = wb.create_sheet(title); s.append(headers)
        for row in data: s.append(row)
        return s
    s = summary
    sheet('摘要与口径', ['项目', '值'], [
        ['数据快照时间', s['captured_at']], ['已完场原始赛事', s['finished_raw']], ['排除特殊赛果', s['excluded']],
        ['比分主样本（每场只计一次）', s['overall']['n']], ['有有效大小球初盘的不同比赛', s['initial_games']],
        ['有已核验大小球赛前末盘的不同比赛', s['closing_games']], ['初盘报价行数（同场可有多家公司）', s['initial_quotes']],
        ['公司数', s['company_count']], ['平均进球', s['overall']['mean_goals']], ['3球及以上比例', s['overall']['over25']],
        ['结算口径', '香港水位；每场每公司每方向各假设投入1单位，盈利=水位、亏损=-1；季度盘拆成相邻两条半球/整数盘各半注。'],
        ['赛果处理', '排除带加时、点球、腰斩、判罚胜等未核实常规时间比分的记录；无特殊备注时使用源站全场比分。'],
        ['临场处理', '仅采用本公司大小球历史中，标记赛前、未封盘、时间严格早于计划开赛的末条；页面l字段不视为已核验临场。'],
        ['盈亏解释', '仅为样本内等额历史结算演示，不是已执行收益，不含费用、限额或可成交性；不得用来声称可持续盈利。'],
        ['区间解释', '比分比例为Wilson 95%区间；ROI为按比赛重采样3000次的95%区间。未处理同队/联赛依赖和选择偏差。'],
        ['公司比较', '不同公司原始覆盖集合不同；同场四公司表使用Bet365/Crown/Sbobet/1xBet共同比赛，盘口可能仍不同。'],
        ['升降盘分组', '分组使用了后续盘口信息，分组初盘盈亏只用于回顾；在开盘时无法预先知道升降方向，不能视为可提前执行的策略。'],
        ['样本限制', '抓取进行中，按日期顺序补采，盘口集中在少数日期；无样本外测试。青年、女子及成年等混合赛事不可直接推广到单一联赛。'],
        ['规则参考', RULES], ['常规时间规则', TIME_RULES]])
    day = sheet('日期与覆盖', ['日期', '比分样本数', '场均进球', '3球及以上比例', '大小球初盘赛事数', '初盘覆盖率', '已核验末盘赛事数'],
                [[r['date'], r['n'], r['mean_goals'], r['over25'], r['initial_games'], r['initial_games']/r['n'], r['closing_games']] for r in s['daily']])
    hist = Counter(g['total'] for g in games)
    dist = sheet('总进球分布', ['总进球', '场数', '占比'], [[g, hist[g], hist[g]/len(games)] for g in range(max(hist)+1)])
    chart = BarChart(); chart.title = '总进球分布（每场计一次）'; chart.x_axis.title='总进球数'; chart.y_axis.title='场数'
    chart.add_data(Reference(dist,min_col=2,min_row=1,max_row=min(dist.max_row,12)),titles_from_data=True)
    chart.set_categories(Reference(dist,min_col=1,min_row=2,max_row=min(dist.max_row,12))); dist.add_chart(chart,'E2')
    sheet('赛事类别', ['类别', '样本数', '场均进球', '3球及以上比例', '4球及以上比例'],
          [[r['category'],r['n'],r['mean_goals'],r['over25'],r['over35']] for r in s['categories']])
    sheet('联赛统计', ['联赛', '比分样本数', '场均进球', '3球及以上比例', '比例95%下限', '比例95%上限', '有初盘比赛数', '样本提示'],
          [[r['league'],r['n'],r['mean_goals'],r['over25'],*r['over25_ci'],r['initial_games'],'可描述，仍需验证' if r['n']>=30 else '不足30场，谨慎解释']
           for r in sorted(s['leagues'],key=lambda r:-r['n'])])
    head=['公司','样本数','大球全赢','大球半赢','走盘','大球半输','大球全输','大球模拟ROI','大球ROI区间下限','大球ROI区间上限','小球模拟ROI','小球ROI区间下限','小球ROI区间上限']
    def book_row(r):
        return [r['company'],r['n'],*[r['over']['counts'].get(k,0) for k in ('全赢','半赢','走盘','半输','全输')],r['over']['roi'],*r['over']['roi_ci'],r['under']['roi'],*r['under']['roi_ci']]
    sheet('各公司初盘结算',head,[book_row(r) for r in s['companies']])
    sheet('同场四公司比较',head,[book_row(r) for r in s['common_four_company']])
    sheet('Bet365按初盘盘口', ['盘口','样本数','大球全赢','大球半赢','走盘','大球半输','大球全输','大球模拟ROI','小球模拟ROI'],
          [[r['line'],r['n'],*[r['over']['counts'].get(k,0) for k in ('全赢','半赢','走盘','半输','全输')],r['over']['roi'],r['under']['roi']] for r in s['reference_initial_lines']])
    sheet('已核验升降盘', ['公司','方向','配对样本数','初盘大球模拟ROI','末盘大球模拟ROI','初盘小球模拟ROI','末盘小球模拟ROI'],
          [[r['company'],r['move'],r['initial']['n'],r['initial']['over']['roi'],r['closing']['over']['roi'],r['initial']['under']['roi'],r['closing']['under']['roi']] for r in s['movements']])
    detail=[]
    for r in rows:
        for phase,label in [('initial','初盘'),('closing','已核验赛前末盘')]:
            q=r[phase]
            if not q: continue
            over,under=settle(r['total'],q['line'],q['over']),settle(r['total'],q['line'],q['under'],'under')
            detail.append([r['id'],r['date'],r['league'],r['home'],r['away'],r['company'],label,r['total'],q['line'],q['over'],q['under'],
                           over['result'],over['profit'],under['result'],under['profit'],r['closing_at'] if phase=='closing' else None,
                           f"{core.BASE}/oddscomp/{r['id']}"])
    sheet('逐场逐庄大小球结算', ['比赛ID','日期','联赛','主队','客队','公司','盘口阶段','总进球','大小球盘口','大球水位','小球水位','大球结算','大球净单位','小球结算','小球净单位','末盘记录时间','来源'], detail)
    sheet('全部赛果及排除原因', ['比赛ID','日期','联赛','主队','客队','主队比分','客队比分','总进球','赛事类别','排除原因','原始备注'],
          [[g['id'],g['date'],g['league'],g['home'],g['away'],g['home_score'],g['away_score'],g['total'],g['category'],g['exclude'],str(g['note'])] for g in snapshot['games']])
    for ws in wb:
        ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions
        for cell in ws[1]:cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='17365D')
        for col in ws.columns:
            ws.column_dimensions[get_column_letter(col[0].column)].width=min(45,max(15,max(len(str(c.value or '')) for c in list(col)[:60])+2))
            header=str(col[0].value)
            if any(x in header for x in ('ROI','比例','占比','覆盖率')):
                for cell in col[1:]:cell.number_format='0.0%'
        for row in ws:
            for cell in row:
                if cell.data_type=='f':cell.data_type='s'
    wb.save(out/'大小球分析.xlsx')


def pct(v):
    return f'{v:.1%}' if v is not None else '无样本'


def write_report(out, s):
    benchmark=next((r for r in s['companies'] if r['cid']==8),None)
    high=sorted([r for r in s['leagues'] if r['n']>=30],key=lambda r:r['over25'],reverse=True)
    lines=[f"# 大小球分析（截至 {s['captured_at']}）",'',
           f"共读取{s['finished_raw']}场已完场赛事；排除{s['excluded']}场常规时间赛果待核实的特殊记录，比分主样本为{s['overall']['n']}场。",
           f"其中只有{s['initial_games']}场有有效大小球初盘，占{pct(s['initial_games']/s['overall']['n'])}；{s['closing_games']}场有已核验的大小球赛前末盘。{s['initial_quotes']}条公司报价对应{s['initial_games']}场比赛，不能按报价行数放大样本量。",'',
           '## 整体进球', '',
           f"场均{s['overall']['mean_goals']:.2f}球；中位数{s['overall']['median_goals']}球。至少2球：{pct(s['overall']['over15'])}；至少3球：{pct(s['overall']['over25'])}；至少4球：{pct(s['overall']['over35'])}。",
           f"3球及以上比例的95%区间为{pct(s['overall']['over25_ci'][0])}至{pct(s['overall']['over25_ci'][1])}。这个比例描述比分，实际大小球输赢还取决于每场盘口和水位。",'',
           '## 各类赛事', '', '|类别|场数|场均进球|3球及以上|','|---|---:|---:|---:|']
    lines += [f"|{r['category']}|{r['n']}|{r['mean_goals']:.2f}|{pct(r['over25'])}|" for r in s['categories']]
    lines += ['', '## 多庄初盘结算', '', '每家公司每场每方向各假设投入1单位，单独统计大球和小球。不同公司覆盖的比赛不同，原始ROI不可直接用于排名。', '',
              '|公司|场数|大球模拟ROI|大球95%区间|小球模拟ROI|', '|---|---:|---:|---|---:|']
    lines += [f"|{r['company']}|{r['n']}|{pct(r['over']['roi'])}|{pct(r['over']['roi_ci'][0])}～{pct(r['over']['roi_ci'][1])}|{pct(r['under']['roi'])}|" for r in s['companies']]
    if benchmark:
        lines += ['',f"以Bet365为单一公司参照：{benchmark['n']}场按实际初盘结算，大球净{benchmark['over']['net_units']:+.2f}单位，小球净{benchmark['under']['net_units']:+.2f}单位。样本内结果不代表未来盈利，当前并未进行样本外验证。"]
    lines += ['', '## 初盘到赛前末盘', '', '只比较同公司同时存在初盘与带时间的赛前末条记录的配对样本。网页即时盘若没有历史时间核验，不进入此表。升降盘分组用到了后续盘口信息，分组初盘盈亏仅用于回顾，不代表开盘时可提前执行的策略。', '',
              '|公司|方向|场数|初盘大球ROI|末盘大球ROI|','|---|---|---:|---:|---:|']
    lines += [f"|{r['company']}|{r['move']}|{r['initial']['n']}|{pct(r['initial']['over']['roi'])}|{pct(r['closing']['over']['roi'])}|" for r in s['movements'] if r['initial']['n']]
    lines += ['', '## 联赛差异（至少30场的描述性样本）', '', '|联赛|场数|场均进球|3球及以上|','|---|---:|---:|---:|']
    selected=high[:4]+[r for r in high[-4:] if r not in high[:4]]
    lines += [f"|{r['league']}|{r['n']}|{r['mean_goals']:.2f}|{pct(r['over25'])}|" for r in selected]
    lines += ['', '## 覆盖偏差与结论', '',
              f"有初盘比赛的3球及以上比例为{pct(s['quoted_games_stats']['over25'])}，未抓到初盘比赛为{pct(s['unquoted_games_stats']['over25'])}；两组的联赛和日期构成可能不同。",'',
              '|日期|比分样本|有初盘|覆盖率|','|---|---:|---:|---:|']
    lines += [f"|{r['date']}|{r['n']}|{r['initial_games']}|{pct(r['initial_games']/r['n'])}|" for r in s['daily']]
    lines += ['', '这份分析可以用于认识已收集赛事的进球分布、检查盘口结算及比较同场不同公司的报价。现阶段没有足够证据据此建立稳定盈利策略；应先补齐盘口覆盖，再按时间划分训练与验证样本。多联赛、多盘口切片属于探索性分析，不把最高收益切片直接视为可交易规律。', '',
              '## 方法与文件', '', '整数盘保留走盘；2.25等季度盘拆成相邻整数和半球盘各半注，分别计算全赢、半赢、走盘、半输和全输。初盘使用各公司f字段；末盘使用大小球历史ou字段，绝不借用亚洲让球历史。',
              f'结算参考：[Goal Line]({RULES})；常规时间参考：[90 Minutes Play]({TIME_RULES})。实际适用规则仍以各公司具体市场规则为准。',
              'input_snapshot.json 保存本次输入快照；大小球分析.xlsx 包含逐场结算、公司比较、联赛统计和排除记录；大小球看板.html 可按公司、日期与盘口阶段筛选。本报告为固定快照，后台继续抓取不会自动改写本次分析。']
    (out/'大小球分析.md').write_text('\n'.join(lines),encoding='utf-8')


def dashboard(out, summary, games, rows):
    data=json.dumps({'summary':summary,'games':games,'rows':rows},ensure_ascii=False).replace('<','\\u003c')
    page='''<!doctype html><html lang="zh-CN"><meta charset="utf-8"><title>大小球分析</title>
<style>body{font:16px system-ui,"Microsoft YaHei",sans-serif;margin:0;background:#f4f6fa;color:#17273b}main{max-width:1200px;margin:36px auto;padding:0 24px}h1{font-size:30px}p{line-height:1.7}.muted{color:#526176}.cards{display:grid;grid-template-columns:repeat(4,1fr);gap:16px}.card,section{background:white;border:1px solid #dbe2ed;border-radius:10px;padding:20px;margin:16px 0}.big{display:block;font-size:29px;color:#174c86;margin:8px 0}select{font:inherit;padding:8px;margin:8px 18px 8px 6px}table{width:100%;border-collapse:collapse;font-size:14px}th,td{padding:9px;text-align:left;border-bottom:1px solid #e8edf3}.scroll{overflow:auto}.bar{height:22px;background:#2a76b8;display:inline-block;vertical-align:middle;margin:5px 8px}.notice{border-left:4px solid #d3982d;background:#fff9ed;padding:14px}a{color:#174c86}@media(max-width:750px){.cards{grid-template-columns:repeat(2,1fr)}}</style>
<main><h1>大小球：赛果与多庄盘口</h1><p id="scope" class="muted"></p><div class="notice">这是已收集样本的历史分析。盘口覆盖不完整；模拟ROI不代表已执行收益或未来盈利。同一场多家报价分别结算，总体进球分布每场仅计一次。</div>
<section><label>公司<select id="book"></select></label><label>日期<select id="day"></select></label><label>盘口阶段<select id="phase"><option value="initial">初盘</option><option value="closing">已核验赛前末盘</option></select></label><p class="muted">以下结算指标只计算筛选后具有对应公司、对应阶段盘口的比赛。</p><div class="cards" id="cards"></div><h3>所选盘口大球结算</h3><div id="grades"></div></section>
<section><h3>进球分布（当前日期的全部比分样本）</h3><p id="histScope" class="muted"></p><div id="hist"></div></section>
<section><h3>逐场核对</h3><p class="muted">展示前200条；完整数据见Excel。水位采用香港格式。</p><div class="scroll"><table><thead><tr><th>日期</th><th>联赛</th><th>对阵</th><th>总进球</th><th>盘口</th><th>大/小水位</th><th>大球结算</th><th>大球净单位</th></tr></thead><tbody id="details"></tbody></table></div></section>
<p><a href="大小球分析.xlsx">完整Excel分析</a> · <a href="大小球分析.md">文字报告</a></p></main><script id="data" type="application/json">__DATA__</script><script>
const D=JSON.parse(document.getElementById('data').textContent);const el=id=>document.getElementById(id);const esc=s=>String(s).replace(/[&<>"']/g,c=>({'&':'&amp;','<':'&lt;','>':'&gt;','"':'&quot;',"'":'&#39;'}[c]));
el('scope').textContent=`快照 ${D.summary.captured_at}｜已完场 ${D.summary.finished_raw} 场｜排除特殊赛果 ${D.summary.excluded} 场｜初盘覆盖 ${D.summary.initial_games} 场`;
const books=[...new Map(D.rows.map(r=>[r.cid,r.company]))];el('book').innerHTML=books.map(([id,n])=>`<option value="${id}">${esc(n)}</option>`).join('');el('book').value=8;
el('day').innerHTML='<option value="all">全部日期</option>'+[...new Set(D.games.map(g=>g.date))].sort().map(d=>`<option>${d}</option>`).join('');
function settle(t,q,side){const parts=[Math.floor(q.line*2)/2,Math.ceil(q.line*2)/2];const signs=parts.map(p=>Math.sign(t-p)*(side==='over'?1:-1));return {grade:signs.reduce((a,b)=>a+b,0),profit:signs.reduce((s,v)=>s+(v>0?q[side]:v<0?-1:0),0)/2};}
const labels={'2':'全赢','1':'半赢','0':'走盘','-1':'半输','-2':'全输'};const pct=v=>Number.isFinite(v)?(v*100).toFixed(1)+'%':'无样本';
function render(){const day=el('day').value,phase=el('phase').value,cid=+el('book').value;const R=D.rows.filter(r=>r.cid===cid&&r[phase]&&(day==='all'||r.date===day));const over=R.map(r=>settle(r.total,r[phase],'over'));const under=R.map(r=>settle(r.total,r[phase],'under'));const mean=a=>a.length?a.reduce((s,v)=>s+v,0)/a.length:NaN;
el('cards').innerHTML=[['可结算比赛数',R.length],['大球模拟ROI',pct(mean(over.map(v=>v.profit)))],['小球模拟ROI',pct(mean(under.map(v=>v.profit)))],['所选样本3球及以上',pct(mean(R.map(r=>r.total>=3?1:0)))]].map(([k,v])=>`<div class="card">${k}<span class="big">${v}</span></div>`).join('');
el('grades').innerHTML=[2,1,0,-1,-2].map(k=>{const n=over.filter(r=>r.grade===k).length;return `<div>${labels[k]} <span class="bar" style="width:${R.length?n/R.length*65:0}%"></span> ${n} 场</div>`}).join('');
const G=D.games.filter(g=>day==='all'||g.date===day);el('histScope').textContent=`${G.length} 场；每场计一次，不受公司选择影响。`;el('hist').innerHTML=Array.from({length:9},(_,i)=>{const n=G.filter(g=>i===8?g.total>=8:g.total===i).length;return `<div>${i===8?'8+':i} 球 <span class="bar" style="width:${G.length?n/G.length*100:0}%"></span> ${n} 场（${pct(n/G.length)}）</div>`}).join('');
el('details').innerHTML=R.slice(0,200).map((r,i)=>`<tr><td>${r.date}</td><td>${esc(r.league)}</td><td>${esc(r.home)} – ${esc(r.away)}</td><td>${r.total}</td><td>${r[phase].line}</td><td>${r[phase].over} / ${r[phase].under}</td><td>${labels[over[i].grade]}</td><td>${over[i].profit.toFixed(2)}</td></tr>`).join('');}
['book','day','phase'].forEach(id=>el(id).addEventListener('change',render));render();</script></html>'''
    (out/'大小球看板.html').write_text(page.replace('__DATA__',data),encoding='utf-8')


if __name__=='__main__':
    parser=argparse.ArgumentParser();parser.add_argument('--refresh',action='store_true')
    parser.add_argument('--out',type=Path,default=core.HERE/'analysis'/'totals_2026-09-09');args=parser.parse_args()
    out=args.out.resolve();out.mkdir(parents=True,exist_ok=True)
    snap=out/'input_snapshot.json'
    snapshot=freeze_snapshot(out) if args.refresh or not snap.exists() else read_json(snap)
    summary,games,rows=analyze(snapshot)
    (out/'summary.json').write_text(json.dumps(summary,ensure_ascii=False,indent=2),encoding='utf-8')
    excel(out,summary,snapshot,games,rows);write_report(out,summary);dashboard(out,summary,games,rows)
    print(json.dumps({k:v for k,v in summary.items() if k not in ('leagues','companies','daily','movements','common_four_company','reference_initial_lines')},ensure_ascii=True))
